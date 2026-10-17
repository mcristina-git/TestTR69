import json
import os
import subprocess
import sys

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

# 1. Recupero documento raw da MongoDB
cmd = [
    "docker", "exec", "-i",
    subprocess.check_output("docker ps -q -f name=_mongo_", shell=True).decode().strip(),
    "mongosh", "--quiet", "--eval",
    "EJSON.stringify(db.getSiblingDB(\"genieacs\").devices.findOne({_id: /BC9424/}))"
]

res = subprocess.run(cmd, capture_output=True, text=True)
if res.returncode != 0 or not res.stdout.strip():
    print("[-] Errore nel recupero dati da MongoDB:", res.stderr)
    sys.exit(1)

raw = res.stdout.strip()
json_idx = raw.find("{")
doc = json.loads(raw[json_idx:])

# 2. Flattening completo di tutti gli 80 parametri
params = {}

def flatten_tree(node, prefix=""):
    if not isinstance(node, dict):
        return
    for k, v in node.items():
        if k.startswith("_"):
            continue
        current_path = f"{prefix}.{k}" if prefix else k
        if isinstance(v, dict):
            if "_value" in v:
                val = v["_value"]
                if isinstance(val, dict):
                    val = next(iter(val.values()))
                params[current_path] = str(val) if val is not None else ""
            elif any(m in v for m in ("_type", "_writable", "_timestamp")):
                params[current_path] = "-"
            else:
                flatten_tree(v, current_path)
        else:
            params[current_path] = str(v)

flatten_tree(doc)
print(f"[+] Totale parametri TR-181 estratti dal data model: {len(params)}")

# 3. Generazione / Aggiornamento cpe_datamodel_flat.xlsx
wb = Workbook()
ws = wb.active
ws.title = "TR-181 Data Model Flat"
ws.views.sheetView[0].showGridLines = True

headers = ["#", "Parameter Path (TR-181 / TR-069)", "Current Value", "Root Object"]
ws.append(headers)

hdr_fill = PatternFill(start_color="0F2537", end_color="0F2537", fill_type="solid")
hdr_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
thin_border = Border(
    left=Side(style="thin", color="D0D7DE"),
    right=Side(style="thin", color="D0D7DE"),
    top=Side(style="thin", color="D0D7DE"),
    bottom=Side(style="thin", color="D0D7DE")
)

for col in range(1, 5):
    c = ws.cell(row=1, column=col)
    c.fill = hdr_fill
    c.font = hdr_font
    c.alignment = Alignment(horizontal="center", vertical="center")

for idx, (path, val) in enumerate(sorted(params.items()), start=1):
    parts = path.split(".")
    root_obj = ".".join(parts[:2]) if len(parts) > 2 else parts[0]
    ws.append([idx, path, val, root_obj])
    
    r_idx = idx + 1
    row_fill = PatternFill(start_color="F7F9FA" if idx % 2 == 0 else "FFFFFF", fill_type="solid")
    for c_idx in range(1, 5):
        c = ws.cell(row=r_idx, column=c_idx)
        c.fill = row_fill
        c.border = thin_border
        c.font = Font(name="Calibri", size=10)
        if c_idx == 1:
            c.alignment = Alignment(horizontal="center")

ws.column_dimensions["A"].width = 8
ws.column_dimensions["B"].width = 65
ws.column_dimensions["C"].width = 40
ws.column_dimensions["D"].width = 25
ws.freeze_panes = "A2"

excel_out = "/home/cristina/cpe-testbed/cpe_datamodel_flat.xlsx"
wb.save(excel_out)
print(f"[+] File Excel aggiornato con tutti i {len(params)} parametri: {excel_out}")

# 4. Aggiornamento report_audit_tr069.pptx
ppt_path = "/home/cristina/cpe-testbed/report_audit_tr069.pptx"

if os.path.exists(ppt_path):
    print(f"[*] Aggiornamento file esistente: {ppt_path}")
    prs = Presentation(ppt_path)
else:
    print(f"[*] Creazione nuova presentazione audit: {ppt_path}")
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

blank_layout = prs.slide_layouts[6]
DARK_BLUE = RGBColor(15, 37, 55)
TEXT_MUTED = RGBColor(100, 116, 139)

items = sorted(params.items())
page_size = 16
pages = (len(items) + page_size - 1) // page_size

for p_num in range(pages):
    chunk = items[p_num * page_size : (p_num + 1) * page_size]
    s = prs.slides.add_slide(blank_layout)
    
    tb = s.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.1))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = f"Audit TR-069 / TR-181 - Parametri Rilevati ({p_num + 1}/{pages})"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = DARK_BLUE
    
    p2 = tf.add_paragraph()
    p2.text = f"CPE TCL 5G - Parametri da {p_num * page_size + 1} a {min((p_num + 1) * page_size, len(items))} di {len(items)}"
    p2.font.size = Pt(12)
    p2.font.color.rgb = TEXT_MUTED
    
    t_shape = s.shapes.add_table(len(chunk) + 1, 2, Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2))
    table = t_shape.table
    table.columns[0].width = Inches(7.5)
    table.columns[1].width = Inches(4.2)
    
    table.cell(0, 0).text = "Parametro TR-181"
    table.cell(0, 1).text = "Valore Registrato"
    for c in range(2):
        cell = table.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = DARK_BLUE
        cell.text_frame.paragraphs[0].font.bold = True
        cell.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
        cell.text_frame.paragraphs[0].font.size = Pt(11)
        
    for r_idx, (path, val) in enumerate(chunk, start=1):
        table.cell(r_idx, 0).text = path
        table.cell(r_idx, 1).text = str(val)[:50]
        for c in range(2):
            cell_p = table.cell(r_idx, c).text_frame.paragraphs[0]
            cell_p.font.name = "Arial"
            cell_p.font.size = Pt(9.5)

prs.save(ppt_path)
print(f"[+] Report audit PowerPoint aggiornato: {ppt_path}")
