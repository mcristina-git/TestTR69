#!/usr/bin/env python3
"""Validator TR-069/TR-181 TCL: TOL, test, modello dati, evidenza PCAP e PPTX.

Eseguire con il Python dell'ambiente virtuale del testbed. I file TOL sorgente,
il payload TR-143 e i risultati restano separati; non sono richiesti altri script.
"""
import os, glob, datetime, openpyxl
import collections
import json
import re
import textwrap
from urllib.parse import urlsplit, urlunsplit
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from copy import copy
import datetime as dt
import fcntl
import hashlib
import math
from pathlib import Path
import signal
import subprocess
import time
import urllib.error
import urllib.parse as url
import urllib.request as http
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.formatting.rule import FormulaRule
import os
import datetime
import shutil
import tempfile
import uuid
import xml.etree.ElementTree as ET

BASE = Path(__file__).resolve().parent

# generate_report_ppt
EXCEL_PATH = str(BASE / 'test.xlsx')
LOGS_DIR = str(BASE / 'logs')
PERIODIC_LOG_PATH = str(BASE / 'periodic_inform_audit.log')
OUTPUT_PPTX = str(BASE / 'report_audit_tr069.pptx')
C_NAVY = RGBColor(16, 37, 66)
C_DARK = RGBColor(33, 37, 41)
C_MUTED = RGBColor(108, 117, 125)
C_LIGHT_BG = RGBColor(245, 247, 250)
C_CARD_BG = RGBColor(255, 255, 255)
C_BORDER = RGBColor(222, 226, 230)
C_PASS = RGBColor(40, 167, 69)
C_FAIL = RGBColor(220, 53, 69)
C_WARN = RGBColor(255, 153, 0)
C_CODE_BG = RGBColor(241, 243, 245)

def apply_background(slide, prs):
    bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg_shape.fill.solid()
    bg_shape.fill.fore_color.rgb = C_LIGHT_BG
    bg_shape.line.fill.background()
    return bg_shape

def add_header(slide, title_text, subtitle_text):
    tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.1))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.name = 'Arial'
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = C_NAVY
    p2 = tf.add_paragraph()
    p2.text = subtitle_text
    p2.font.name = 'Arial'
    p2.font.size = Pt(11)
    p2.font.color.rgb = C_MUTED
    p2.space_before = Pt(4)

def add_card(slide, left, top, width, height, bg_color=C_CARD_BG, border_color=C_BORDER):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    if border_color:
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
    else:
        card.line.fill.background()
    return card
SENSITIVE_PATH = re.compile('password|passwd|passphrase|secret|token|private.?key|presharedkey|wepkey|username|credential|authorization|pincode|pukcode|(^|[._])(?:key|pin|puk|psk)(?:[._]|$)', re.IGNORECASE)

def sanitise_text(text):
    """Oscura righe contenenti credenziali anche negli estratti di log."""
    return '\n'.join(('[riga con dati sensibili oscurata]' if SENSITIVE_PATH.search(line) else line for line in str(text).splitlines()))

def datamodel_parameters(device):
    """Visita anche gli oggetti con metadati, senza confonderli con foglie."""
    parameters = []

    def walk(node, path):
        if not isinstance(node, dict):
            return
        if node.get('_object') is False or '_value' in node:
            parameters.append((path, node))
        for name, child in node.items():
            if not name.startswith('_'):
                walk(child, f'{path}.{name}')
    for root in ('Device', 'InternetGatewayDevice'):
        if root in device:
            walk(device[root], root)
    return sorted(parameters)

def value_state(node):
    if '_value' not in node:
        return 'non letto'
    if node['_value'] is None:
        return 'null'
    if node['_value'] == '':
        return 'vuoto'
    return 'valorizzato'

def display_value(path, node):
    state = value_state(node)
    if state == 'non letto':
        return '[non ancora letto]'
    if SENSITIVE_PATH.search(path):
        return '[oscurato]'
    if state == 'null':
        return '[null]'
    if state == 'vuoto':
        return '[stringa vuota]'
    value = node['_value']
    if isinstance(value, str):
        if '://' in value:
            try:
                parsed = urlsplit(value)
                value = urlunsplit((parsed.scheme, parsed.netloc.rsplit('@', 1)[-1], '/[oscurato]' if 'ConnectionRequestURL' in path else parsed.path, '', ''))
            except ValueError:
                return '[URL oscurato]'
        return re.sub('[\\x00-\\x08\\x0b\\x0c\\x0e-\\x1f]', lambda m: repr(m[0])[1:-1], value)
    return json.dumps(value, ensure_ascii=False, default=str)

def credential_values(parameters):
    """Protegge anche credenziali ripetute come parte di altri campi."""
    return sorted({str(node['_value']) for path, node in parameters if SENSITIVE_PATH.search(path) and isinstance(node.get('_value'), str) and (len(node['_value']) >= 8)}, key=len, reverse=True)

def redact_secrets(text, secrets):
    for secret in secrets:
        text = str(text).replace(secret, '[oscurato]')
    return text

def wrapped_lines(text, width):
    lines = []
    for line in str(text).split('\n'):
        lines.extend(textwrap.wrap(line, width=width, expand_tabs=False, replace_whitespace=False, drop_whitespace=False, break_on_hyphens=False) or [''])
    return lines

def set_table_cell(cell, text, header=False, font_size=11):
    cell.text = str(text)
    cell.margin_left = cell.margin_right = Inches(0.07)
    cell.margin_top = cell.margin_bottom = Inches(0.045)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.text_frame.word_wrap = True
    cell.fill.solid()
    cell.fill.fore_color.rgb = C_NAVY if header else C_CARD_BG
    for paragraph in cell.text_frame.paragraphs:
        paragraph.font.name = 'Arial'
        paragraph.font.size = Pt(font_size)
        paragraph.font.bold = header
        paragraph.font.color.rgb = RGBColor(255, 255, 255) if header else C_DARK
        paragraph.space_before = paragraph.space_after = Pt(0)
        paragraph.line_spacing = Pt(14)

def add_datamodel_summary(prs, parameters, device):
    counts = collections.Counter((value_state(n) for _, n in parameters))
    version = next((n.get('_value') for p, n in parameters if p == 'Device.RootDataModelVersion'), 'non disponibile')
    branches = collections.defaultdict(collections.Counter)
    for path, node in parameters:
        branches['.'.join(path.split('.')[:2])][value_state(node)] += 1
    branch_items = sorted(branches.items())
    pages = [branch_items[i:i + 12] for i in range(0, len(branch_items), 12)]
    for page_number, entries in enumerate(pages, 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        apply_background(slide, prs)
        add_header(slide, f'2. Data model completo ({page_number}/{len(pages)})', f"Fonte: cache ACS del CPE. Device:{version}. Ultimo Inform: {device.get('_lastInform', 'non disponibile')}")
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(1.45), Inches(11.73), Inches(0.45))
        paragraph = tb.text_frame.paragraphs[0]
        paragraph.text = f"{len(parameters)} parametri: {counts['valorizzato']} non vuoti, {counts['vuoto']} vuoti, {counts['null']} null, {counts['non letto']} non letti"
        paragraph.font.size = Pt(16)
        paragraph.font.color.rgb = C_NAVY
        table = slide.shapes.add_table(len(entries) + 1, 4, Inches(0.8), Inches(2.0), Inches(11.73), Inches((len(entries) + 1) * 0.35)).table
        for col, width in zip(table.columns, (6.7, 1.6, 1.6, 1.83)):
            col.width = Inches(width)
        rows = [('Ramo', 'Parametri', 'Vuoti / null', 'Non letti')]
        rows += [(branch, sum(c.values()), c['vuoto'] + c['null'], c['non letto']) for branch, c in entries]
        for i, row in enumerate(rows):
            for j, value in enumerate(row):
                set_table_cell(table.cell(i, j), value, header=i == 0)
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(6.92), Inches(11.73), Inches(0.35))
        p = tb.text_frame.paragraphs[0]
        p.text = "Elenco completo nelle slide successive all'audit. Credenziali oscurate. Un valore presente non implica conformit?."
        p.font.size = Pt(10)
        p.font.color.rgb = C_MUTED

def add_datamodel_details(prs, parameters):
    """Paginazione per ramo e altezza: nessun limite al numero di parametri."""
    secrets = credential_values(parameters)
    groups = collections.defaultdict(list)
    for path, node in parameters:
        groups['.'.join(path.split('.')[:2])].append((path, node))
    slide_count = 0
    for branch, items in sorted(groups.items()):
        pages, page, height = ([], [], 0.0)
        for path, node in items:
            path_lines = wrapped_lines(path, 69)
            value_lines = wrapped_lines(redact_secrets(display_value(path, node), secrets), 39)
            meta = [str(node.get('_type', 'tipo n/d')).removeprefix('xsd:'), 'R/W' if node.get('_writable') is True else 'R/O' if node.get('_writable') is False else 'accesso n/d']
            if value_state(node) != 'valorizzato':
                meta.append(value_state(node))
            chunks = [value_lines[i:i + 20] for i in range(0, len(value_lines), 20)]
            for part, chunk in enumerate(chunks, 1):
                shown_path = path_lines + ([f'(continua {part}/{len(chunks)})'] if len(chunks) > 1 else [])
                row_height = max(len(shown_path), len(chunk), len(meta)) * 14 / 72 + 0.12
                if page and height + row_height > 4.85:
                    pages.append(page)
                    page, height = ([], 0.0)
                page.append(('\n'.join(shown_path), '\n'.join(chunk), '\n'.join(meta), row_height))
                height += row_height
        if page:
            pages.append(page)
        for number, rows in enumerate(pages, 1):
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            apply_background(slide, prs)
            add_header(slide, 'Data model completo', f'{branch}   ({number}/{len(pages)})')
            table = slide.shapes.add_table(len(rows) + 1, 3, Inches(0.8), Inches(1.6), Inches(11.73), Inches(0.36 + sum((r[3] for r in rows)))).table
            for col, width in zip(table.columns, (6.5, 3.9, 1.33)):
                col.width = Inches(width)
            table.rows[0].height = Inches(0.36)
            for j, title in enumerate(('Parametro', 'Valore', 'Tipo / accesso')):
                set_table_cell(table.cell(0, j), title, header=True)
            for i, row in enumerate(rows, 1):
                table.rows[i].height = Inches(row[3])
                for j in range(3):
                    set_table_cell(table.cell(i, j), row[j])
            tb = slide.shapes.add_textbox(Inches(0.8), Inches(6.95), Inches(11.73), Inches(0.28))
            p = tb.text_frame.paragraphs[0]
            p.text = f'Fonte: ACS. Parametri del ramo: {len(items)}. R/O: sola lettura. R/W: scrivibile. Credenziali oscurate.'
            p.font.size = Pt(9)
            p.font.color.rgb = C_MUTED
            slide_count += 1
    return slide_count

# Il modello contiene le slide editoriali 9-13, preparate fuori dal validator.
REPORT_TEMPLATE = BASE / 'report_template_tr069.pptx'


def template_shape(slide, shape_id):
    for shape in slide.shapes:
        if shape.shape_id == shape_id:
            return shape
    raise TestBlocked('Modello PPTX incompatibile: oggetto mancante ' + str(shape_id))


def replace_paragraph(paragraph, text):
    """Sostituisce il contenuto mantenendo font, colori e stile del modello."""
    from copy import deepcopy
    style = deepcopy(paragraph.runs[0]._r.rPr) if paragraph.runs and paragraph.runs[0]._r.rPr is not None else None
    paragraph.text = str(text)
    if style is not None and paragraph.runs:
        run = paragraph.runs[0]._r
        if run.rPr is not None:
            run.remove(run.rPr)
        run.insert(0, style)


def template_text(shape, lines):
    """Riusa anche gli stili dei paragrafi successivi e le distanze originali."""
    from copy import deepcopy
    tf = shape.text_frame
    old = [deepcopy(p._p) for p in tf.paragraphs]
    tf.clear()
    for index, line in enumerate(lines):
        para_xml = deepcopy(old[min(index, len(old)-1)])
        if index == 0:
            tf._txBody.replace(tf.paragraphs[0]._p, para_xml)
        else:
            tf._txBody.append(para_xml)
        replace_paragraph(tf.paragraphs[index], line)


def remove_embedded_workbooks(prs):
    """Rimuove oggetti OLE e relazioni: nessun Excel incorporato nel pacchetto."""
    for slide in prs.slides:
        for shape in list(slide.shapes):
            if shape._element.xpath('.//p:oleObj'):
                element = shape._element
                element.getparent().remove(element)
        used = set()
        for element in slide._element.iter():
            used.update(value for key, value in element.attrib.items()
                        if key.startswith('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'))
        for rel in list(slide.part.rels.values()):
            if rel.rId not in used and rel.reltype.rsplit('/', 1)[-1] in ('oleObject', 'package', 'image'):
                slide.part.drop_rel(rel.rId)


def report_tol_snapshot(device):
    """Legge gli esiti reali, con controllo di identita e sorgente del workbook."""
    if not OUTPUT.is_file():
        raise TestBlocked('TOL_TCL.xlsx assente: eseguire la suite TOL')
    wb = openpyxl.load_workbook(OUTPUT, data_only=True)
    expected = hashlib.sha256(device['_id'].encode()).hexdigest()
    if wb.properties.identifier != expected:
        raise TestBlocked('La TOL appartiene a un altro CPE')
    original = openpyxl.load_workbook(SOURCE, data_only=True)['MASTER']
    ws = wb['MASTER']
    rows = [r for r in range(2, original.max_row+1) if original.cell(r, 5).value]
    for r in rows:
        for c in range(1, 10):
            if ws.cell(r, c).value != original.cell(r, c).value:
                raise TestBlocked('La TOL non corrisponde al modello originale A-I')
    counts = collections.Counter(str(ws.cell(r, 10).value or 'NON ESEGUITO').split('\n')[0] for r in rows)
    manual = {r[0] for r in wb['Valori_TCL'].iter_rows(min_row=2, values_only=True)
              if r[2] == '[NOTE]' and r[5] == 'VALIDAZIONE UTENTE'}
    meta = json.loads((BASE/'logs'/'tol_tcl_latest.json').read_text())
    if meta.get('source_sha256') != hashlib.sha256(SOURCE.read_bytes()).hexdigest():
        raise TestBlocked('Sorgente TOL modificato: rigenerare la TOL')
    return len(rows), counts, len(manual), meta.get('generated_at', 'n/d')


def update_template_tol(slide, device):
    table_shape = template_shape(slide, 7)
    table_shape.left, table_shape.top = Inches(.8), Inches(1.6)
    table_shape.width, table_shape.height = Inches(7.0), Inches(3.6)
    table = table_shape.table
    from copy import deepcopy
    while len(table.rows) < 7:
        table._tbl.append(deepcopy(table.rows[len(table.rows)-1]._tr))
    table.columns[0].width, table.columns[1].width = Inches(5.4), Inches(1.6)
    try:
        total, counts, manual, generated = report_tol_snapshot(device)
        data = [('Esito TOL', 'Casi'), ('Totale casi di test', total),
                ('PASS', counts['PASS']), ('PARZIALE', counts['PARZIALE']),
                ('FAIL', counts['FAIL']), ('BLOCCATO', counts['BLOCCATO']),
                ('NON ESEGUITO', counts['NON ESEGUITO'])]
        notes = [f'{manual} esiti includono una validazione esplicita dell’utente. Le note Excel conservano l’esito automatico.',
                 f'TOL aggiornata: {generated}. Dettagli e valori nel file TOL_TCL.xlsx, nella cartella del progetto.']
    except (TestBlocked, OSError, KeyError, ValueError) as error:
        data = [('Esito TOL', 'Casi'), ('Totale casi di test', 'n/d')] + [(x, 'n/d') for x in ('PASS','PARZIALE','FAIL','BLOCCATO','NON ESEGUITO')]
        notes = ['Riepilogo non disponibile: '+str(error), 'Eseguire la suite TOL per aggiornare il file Excel.']
    for i, values in enumerate(data):
        table.rows[i].height = Inches(.49)
        for j, value in enumerate(values):
            set_table_cell(table.cell(i,j), value, header=i==0, font_size=14)
    note = template_shape(slide, 4)
    note.left, note.top, note.width, note.height = Inches(.8), Inches(5.65), Inches(11.7), Inches(1.1)
    template_text(note, notes)
    for paragraph in note.text_frame.paragraphs:
        paragraph.font.size = Pt(12)
        paragraph.font.name = 'Arial'
        paragraph.font.color.rgb = C_MUTED


def update_template_model(prs, parameters, device):
    counts = collections.Counter(value_state(n) for _,n in parameters)
    version = next((n.get('_value') for p,n in parameters if p=='Device.RootDataModelVersion'), 'n/d')
    branches = collections.defaultdict(collections.Counter)
    for path,node in parameters:
        branches['.'.join(path.split('.')[:2])][value_state(node)] += 1
    items = sorted(branches.items())
    # Due pagine nel corpo del modello, dettagli integrali dopo le slide editoriali.
    cut = max(12, math.ceil(len(items)/2))
    pages = [items[:cut], items[cut:]]
    from copy import deepcopy
    for index, entries in enumerate(pages):
        slide = prs.slides[3+index]
        template_text(template_shape(slide,3), [f'2. Data model completo ({index+1}/2)',
            f"Fonte: cache ACS del CPE. Device:{version}. Ultimo Inform: {device.get('_lastInform','n/d')}"])
        template_text(template_shape(slide,4), [f"{len(parameters)} parametri: {counts['valorizzato']} non vuoti, {counts['vuoto']} vuoti, {counts['null']} null, {counts['non letto']} non letti"])
        shape=template_shape(slide,5); table=shape.table
        required=len(entries)+1
        while len(table.rows)<required:
            table._tbl.append(deepcopy(table.rows[len(table.rows)-1]._tr))
        while len(table.rows)>required:
            table._tbl.remove(table.rows[len(table.rows)-1]._tr)
        row_height=min(.35, 4.65/max(required,1))
        rows=[('Ramo','Parametri','Vuoti / null','Non letti')]+[(b,sum(c.values()),c['vuoto']+c['null'],c['non letto']) for b,c in entries]
        for i,row in enumerate(rows):
            table.rows[i].height=Inches(row_height)
            for j,value in enumerate(row):set_table_cell(table.cell(i,j),value,header=i==0)
        template_text(template_shape(slide,6), ['Elenco completo in appendice, dopo la slide 13. Excel nella cartella. Credenziali oscurate. Un valore presente non implica conformità.'])


def update_template_periodic(prs, device, secrets):
    interval=get_param_value(device,INTERVAL)
    try:
        proof=verified_periodic_evidence(device,int(interval or 0))
        status='PASS' if proof['passed'] else 'FAIL'
    except (TestBlocked,ValueError,KeyError,TypeError,OSError,subprocess.SubprocessError) as error:
        proof=None;status='WARN';reason=str(error)
    metrics=['Metriche Periodicità','Parametro monitorato',INTERVAL,'Intervallo corrente ACS',f'{interval} secondi',
             'Campioni e tolleranza','3 Inform, 2 intervalli, tolleranza ±5 s','Esito della verifica',status,
             'Fonte evidenza','PCAP verificato e correlazione con il log ACS']
    template_text(template_shape(prs.slides[6],5), metrics)
    log=Path(PERIODIC_LOG_PATH).read_text() if Path(PERIODIC_LOG_PATH).exists() else 'Log non disponibile.'
    template_text(template_shape(prs.slides[6],7), ['Log verifica Inform (periodic_inform_audit.log)',redact_secrets(sanitise_text(log.strip()),secrets)])
    table=template_shape(prs.slides[7],4).table
    rows=[['#','Timestamp UTC','Frame req/risp','TCP stream','Delta (s)','Evento / risposta']]
    if proof:
        for i,row in enumerate(proof['informs']):
            rows.append([str(i+1),row['timestamp'],f"{row['frame']} / {row['response_frame']}",str(row['stream']),
                         '-' if i==0 else f"{proof['intervals_s'][i-1]:.6f}",f"{row['event_codes']}\nHTTP {row['response_http']}; retry {row['retry_count']}"])
        entries=[status+f": intervallo atteso {proof['expected_interval_s']} s, tolleranza ±5 s.",
                 'PCAP: '+proof['pcap'],'SHA256: '+proof['pcap_sha256'],
                 'Cattura: '+proof['started']+' - '+proof['ended'],
                 'Tre Inform e tre InformResponse. I numeri frame si riferiscono al PCAP estratto.']
    else:
        rows += [[str(i+1),'n/d','n/d','n/d','n/d','Non verificato'] for i in range(3)]
        entries=['WARN: '+reason,'Eseguire lo script senza --report-only per acquisire una prova aggiornata.']
    for i,row in enumerate(rows):
        for j,value in enumerate(row):set_table_cell(table.cell(i,j),value,header=i==0,font_size=11)
    note=template_shape(prs.slides[7],5)
    template_text(note,entries)
    for i,paragraph in enumerate(note.text_frame.paragraphs):
        paragraph.font.name='Arial';paragraph.font.size=Pt(14 if i==0 else 11)
        paragraph.font.color.rgb=({'PASS':C_PASS,'FAIL':C_FAIL,'WARN':C_WARN}[status] if i==0 else C_DARK)


def build_presentation(device=None):
    """Aggiorna il modello caricato, conserva 9-13, aggiunge il modello dati integrale."""
    if device is None:device=fetch_device()
    if not device:raise TestBlocked('Nessun CPE disponibile per il report')
    parameters=datamodel_parameters(device)
    if not parameters:raise TestBlocked('Nessun parametro del data model disponibile')
    template=Path(REPORT_TEMPLATE)
    if not template.is_file():raise TestBlocked('Modello PPTX assente: '+str(template))
    prs=Presentation(str(template))
    if len(prs.slides)!=13:raise TestBlocked('Il modello deve contenere le 13 slide di base, senza appendice generata')
    remove_embedded_workbooks(prs)
    secrets=credential_values(parameters)
    # Cambia solo la data, conservando i run formattati della copertina e l’autrice.
    cover=template_shape(prs.slides[0],3)
    for paragraph in cover.text_frame.paragraphs:
        for run in paragraph.runs:
            run.text=re.sub(r'(?<=Data: )\d{2}/\d{2}/\d{4} \d{2}:\d{2}:\d{2}',
                            datetime.datetime.now().strftime('%d/%m/%Y %H:%M:%S'),run.text)
    update_template_tol(prs.slides[2],device)
    update_template_model(prs,parameters,device)
    # Il log SetParameterValues resta storico: nessun SET per generare un report.
    template_text(template_shape(prs.slides[5],5), ['Dettagli Modifica Parametro','Operazione documentata',
        'SetParameterValues','Fonte','Log SOAP archiviato, prova storica a 120 s',
        'Configurazione corrente',f'PeriodicInformInterval: {get_param_value(device,INTERVAL)} s',
        'Generazione report','La sola generazione del PPTX non modifica i parametri'])
    update_template_periodic(prs,device,secrets)
    detail_slides=add_datamodel_details(prs,parameters)
    from zipfile import ZipFile
    # Salvataggio atomico: non sovrascrivere il report se il pacchetto contiene allegati.
    temporary=Path(OUTPUT_PPTX).with_suffix('.tmp.pptx')
    try:
        prs.save(str(temporary))
        with ZipFile(temporary) as archive:
            if any(n.startswith('ppt/embeddings/') for n in archive.namelist()):
                raise TestBlocked('Il PPTX contiene ancora file incorporati')
        os.replace(temporary,OUTPUT_PPTX)
    finally:
        temporary.unlink(missing_ok=True)
    print(f'[+] Modello {template.name}: 13 slide base e {detail_slides} slide con tutti i {len(parameters)} parametri')
    print('[+] Slide editoriali 9-13 conservate dal modello. Excel separati, nessun allegato nel PPTX.')
    return OUTPUT_PPTX

# tol_tcl
SOURCE = BASE / 'tol_standard.xlsx'
OUTPUT = BASE / 'TOL_TCL.xlsx'
NBI = os.environ.get('GENIEACS_NBI_URL', 'http://127.0.0.1:7547')
INTERVAL = 'Device.ManagementServer.PeriodicInformInterval'
PIT = 'Device.ManagementServer.PeriodicInformTime'
DIRECT = {8, 9, 10, 11, 15, 18, 20, 30, 31, 32, 39, 40, 41, 42, 43, 44, 45, 46, 47, 48, 49, 50, 51, 52, 53, 56, 57, 58, 59, 62, 63, 64, 67, 74}
GET_ROWS = {7, 9, 18, 26, 31, 32, 33, 34, 35, 38, 39, *range(40, 59), 61, 62, 63, 64, 65, 66, 67, 68, 70, 71, 72, 73, 74, 75, 76}
ACTIVE_ROWS = {8, 10, 11, 15, 20, 30, 57, 59}

def utcnow():
    return dt.datetime.now(dt.timezone.utc)

def parsedate(value):
    return dt.datetime.fromisoformat(str(value).replace('Z', '+00:00'))

def counter_value(node):
    value = node.get('_value')
    kind = node.get('_type')
    limits = {'xsd:unsignedInt': 2 ** 32, 'xsd:unsignedLong': 2 ** 64}
    if kind not in limits or type(value) not in (int, str):
        return None
    if isinstance(value, str) and (not re.fullmatch('[+]?[0-9]+', value)):
        return None
    try:
        number = int(value)
    except (ValueError, TypeError):
        return None
    return number if 0 <= number < limits[kind] else None

class TestBlocked(RuntimeError):
    pass

class CWMPFault(RuntimeError):
    pass

class ACS:

    def __init__(self, folder):
        self.folder = folder
        self.tasks = []
        self.restore_tasks = set()
        self.protection = {}
        self.restore_failed = False
        self.device_id = None
        _, rows = self.api('/devices', {'projection': '_id'})
        if len(rows) != 1:
            raise TestBlocked('Selezione CPE non univoca')
        self.device_id = rows[0]['_id']
        _, pending = self.api('/tasks', {'query': json.dumps({'device': self.device_id}), 'projection': '_id,name'})
        if pending:
            raise TestBlocked('Coda ACS preesistente non vuota')
        self.acs = subprocess.check_output(['docker', 'ps', '-q', '-f', 'name=_acs_'], text=True, timeout=10).strip()
        if not self.acs or '\n' in self.acs:
            raise TestBlocked('Container ACS non univoco')
        self.device = self.fetch()
        self.secrets = credential_values(datamodel_parameters(self.device))

    def event(self, kind, **data):
        line = redact_secrets(json.dumps({'time': utcnow().isoformat(), 'event': kind, **data}, default=str), self.secrets)
        with (self.folder / 'events.jsonl').open('a') as f:
            f.write(line + '\n')
        print(line, flush=True)

    def api(self, path, params=None, data=None, method=None):
        target = NBI + path + ('?' + url.urlencode(params) if params else '')
        request = http.Request(target, data=json.dumps(data).encode() if data is not None else None, headers={'Content-Type': 'application/json'}, method=method)
        try:
            with http.urlopen(request, timeout=15) as response:
                raw = response.read()
                return (response.status, json.loads(raw) if raw else None)
        except urllib.error.HTTPError as error:
            raise RuntimeError(f"ACS HTTP {error.code} su {path.split('?')[0].split('/devices/')[0]}") from None

    def fetch(self):
        _, rows = self.api('/devices', {'query': json.dumps({'_id': self.device_id})})
        if len(rows) != 1:
            raise TestBlocked('CPE non disponibile')
        self.device = rows[0]
        return self.device

    def node(self, path):
        current = self.device
        for part in path.split('.'):
            if not isinstance(current, dict) or part not in current:
                return {}
            current = current[part]
        return current if isinstance(current, dict) else {}

    def value(self, path):
        return self.node(path).get('_value')

    def post(self, task, restore=False):
        status, row = self.api('/devices/' + url.quote(self.device_id, safe='') + '/tasks', data=task)
        task_id = row['_id']
        self.tasks.append(task_id)
        if restore:
            self.restore_tasks.add(task_id)
        self.event('task', name=task['name'], task_id=task_id, http=status, restore=restore)
        return task_id

    def wait(self, task_ids, seconds=100):
        pending = set(task_ids)
        end = time.monotonic() + seconds
        while pending and time.monotonic() < end:
            for task_id in list(pending):
                _, faults = self.api('/faults', {'query': json.dumps({'_id': self.device_id + ':task_' + task_id}), 'projection': 'code,message,timestamp'})
                if faults:
                    codes = [f.get('code', '?') for f in faults]
                    self.event('fault', task_id=task_id, codes=codes)
                    raise CWMPFault('Fault ACS/CWMP: ' + ', '.join(codes))
                _, rows = self.api('/tasks', {'query': json.dumps({'_id': task_id}), 'projection': '_id'})
                if not rows:
                    pending.remove(task_id)
            if pending:
                time.sleep(2)
        if pending:
            raise TimeoutError('Task non eseguito entro il limite')
        self.fetch()

    def refresh(self, paths=None, attributes=False):
        start = utcnow()
        if attributes:
            task = {'name': 'provisions', 'provisions': [['refresh', p, 1, False, 'notification'] for p in paths]}
        elif paths:
            task = {'name': 'getParameterValues', 'parameterNames': list(paths)}
        else:
            task = {'name': 'refreshObject', 'objectName': 'Device'}
        self.wait([self.post(task)])
        if paths:
            for path in paths:
                stamp = self.node(path).get('_attributesTimestamp' if attributes else '_timestamp')
                if not stamp or parsedate(stamp) < start - dt.timedelta(seconds=1):
                    raise TestBlocked('Risposta non aggiornata: ' + path)
        else:
            self.refreshed_at = start
        return self.device

    def change(self, changes, restore=False):
        if self.restore_failed and (not restore):
            raise TestBlocked('Ulteriori scritture sospese: ripristino non verificato')
        parameters = []
        for path, value in changes.items():
            node = self.node(path)
            if not node.get('_writable'):
                raise TestBlocked('Parametro non scrivibile: ' + path)
            if '_value' not in node:
                raise TestBlocked('Valore iniziale non disponibile: ' + path)
            if not restore:
                self.protection.setdefault(path, (node['_value'], node.get('_type', 'xsd:string')))
            parameters.append([path, value, node.get('_type', 'xsd:string')])
        task = self.post({'name': 'setParameterValues', 'parameterValues': parameters}, restore=restore)
        self.wait([task])
        start = utcnow()
        read = self.post({'name': 'getParameterValues', 'parameterNames': list(changes)}, restore=restore)
        self.wait([read])
        for path, wanted in changes.items():
            node = self.node(path)
            operational = path.endswith('.DiagnosticsState') and wanted == 'Requested'
            actual = node.get('_value')
            equal = actual == wanted
            if node.get('_type') == 'xsd:dateTime':
                try:
                    equal = parsedate(actual) == parsedate(wanted)
                except (TypeError, ValueError):
                    pass
            if not operational and (not equal):
                raise CWMPFault('Rilettura diversa dal valore richiesto: ' + path)
            if not node.get('_timestamp') or parsedate(node['_timestamp']) < start - dt.timedelta(seconds=1):
                raise TestBlocked('Rilettura non aggiornata: ' + path)
        self.event('change_verified', paths=list(changes), restore=restore)

    def restore(self, paths):
        changes = {p: self.protection[p][0] for p in paths if p in self.protection}
        if not changes:
            return
        try:
            self.change(changes, restore=True)
            for path in changes:
                self.protection.pop(path, None)
            if not self.protection:
                self.restore_failed = False
        except Exception:
            self.restore_failed = True
            self.event('restore_failed', paths=list(changes))
            raise

    def lines(self):
        result = subprocess.run(['docker', 'exec', self.acs, 'tail', '-n', '1200', '/var/log/cwmp.log'], capture_output=True, text=True, timeout=10, check=True)
        return result.stdout.splitlines()

    def periodic(self, since):
        rows = []
        for line in self.lines():
            if self.device_id not in line or 'Inform;' not in line:
                continue
            timestamp = parsedate(line.split()[0])
            code = re.search('informEvent="([^"]*)"', line)
            retry = re.search('informRetryCount=(\\d+)', line)
            request = re.search('cpeRequestId="([^"]*)"', line)
            if timestamp >= since and code and ('2 PERIODIC' in [x.strip() for x in code[1].split(',')]):
                rows.append({'timestamp': timestamp.isoformat(), 'event_codes': code[1], 'retry_count': int(retry[1]) if retry else None, 'request_id': request[1] if request else None})
        return rows

    def observe(self, interval, cycles, since=None):
        since = since or utcnow()
        deadline = time.monotonic() + interval * (cycles + 2) + 40
        heartbeat = 0
        while time.monotonic() < deadline:
            rows = self.periodic(since)
            if len(rows) >= cycles + 1:
                rows = rows[:cycles + 1]
                times = [parsedate(r['timestamp']) for r in rows]
                deltas = [(b - a).total_seconds() for a, b in zip(times, times[1:])]
                passed = all((abs(x - interval) <= 5 for x in deltas)) and all((r['retry_count'] == 0 for r in rows))
                return {'informs': rows, 'intervals_s': deltas, 'passed': passed, 'expected_interval_s': interval, 'tolerance_s': 5}
            if time.monotonic() >= heartbeat:
                self.event('observing', interval_s=interval, cycles=cycles, received=len(rows), required=cycles + 1)
                heartbeat = time.monotonic() + 20
            time.sleep(2)
        raise TimeoutError(f'Osservazione incompleta: attesi {cycles + 1} Inform periodici')

    def cleanup(self):
        for task_id in self.tasks:
            try:
                _, rows = self.api('/tasks', {'query': json.dumps({'_id': task_id}), 'projection': '_id'})
                if rows and task_id not in self.restore_tasks:
                    self.api('/tasks/' + url.quote(task_id, safe=''), method='DELETE')
                _, faults = self.api('/faults', {'query': json.dumps({'_id': self.device_id + ':task_' + task_id}), 'projection': '_id'})
                if task_id not in self.restore_tasks:
                    for fault in faults:
                        self.api('/faults/' + url.quote(fault['_id'], safe=''), method='DELETE')
            except Exception:
                self.event('cleanup_incomplete', task_id=task_id)

class TOLSuite:

    def __init__(self, acs, active=False):
        self.acs = acs
        self.active = active
        self.results = {}
        self.full_model = {}
        self.exported = False
        self.source_hash = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
        self.manual_reviews = previous_manual_reviews(getattr(acs,'device_id',''),self.source_hash)
        self.workbook = openpyxl.load_workbook(SOURCE)
        self.sheet = self.workbook['MASTER']
        self.rows = [r for r in range(2, self.sheet.max_row + 1) if self.sheet.cell(r, 5).value]
        if self.rows != list(range(2, 80)):
            raise TestBlocked('TOL sorgente cambiato: verificare il mapping delle righe prima dei test')
        for row, expected in {8: 'CPE_0004_tr069', 10: 'CPE_0006_tr069', 11: 'CPE_0052_tr069', 15: 'CPE_0047_tr069', 20: 'CPE_0021_tr069', 30: 'CPE_00075_tr069', 57: 'CPE_00079_tr069', 59: 'CPE_0054_tr069'}.items():
            if self.sheet.cell(row, 2).value != expected:
                raise TestBlocked('Mapping TOL cambiato alla riga ' + str(row))
        for row in self.rows:
            self.results[row] = {'status': 'NON ESEGUITO', 'values': {}, 'notes': 'Fuori dal perimetro GET e dei 34 test diretti selezionati.', 'scope': 'nessuno'}
        for row in ACTIVE_ROWS:
            self.results[row]['notes'] = 'Prova attiva disponibile con --tol-active; richiede variazione reale, rilettura e ripristino.'

    @property
    def parameters(self):
        return dict(datamodel_parameters(self.acs.device))

    def result(self, row, status, values=None, notes='', scope='GET'):
        safe = {}
        for path, node in (values or {}).items():
            safe[path] = redact_secrets(display_value(path, node) if isinstance(node, dict) else str(node), self.acs.secrets)
        self.results[row] = {'status': status, 'values': safe, 'notes': redact_secrets(str(notes), self.acs.secrets), 'scope': scope}
        self.acs.event('test_result', row=row, status=status, parameters=len(safe), scope=scope)

    def matched(self, prefixes=(), names=()):
        return {p: n for p, n in self.parameters.items() if p in names or any((p.startswith(x) for x in prefixes))}

    def get(self, row, prefixes=(), names=(), partial=False, notes='', nonempty=()):
        values = self.matched(prefixes, names)
        missing = [p for p in names if p not in values]
        unread = [p for p, n in values.items() if '_value' not in n]
        empty_required = [p for p in nonempty if p in values and values[p].get('_value') in ('', None)]
        if not values:
            return self.result(row, 'BLOCCATO', notes='Nessun parametro corrispondente osservato nel modello TCL. ' + notes)
        if missing or unread or empty_required:
            reason = '; '.join(['assenti: ' + ', '.join(missing)] if missing else [])
            if unread:
                reason += '; non letti: ' + ', '.join(unread)
            if empty_required:
                reason += '; vuoti richiesti: ' + ', '.join(empty_required)
            return self.result(row, 'PARZIALE', values, reason + '. ' + notes)
        self.result(row, 'PARZIALE' if partial else 'PASS', values, notes or 'GPV aggiornato; esito relativo alla lettura, non a prove funzionali ulteriori.')

    def read_tests(self, first_sample):
        a = self.acs
        ntp_servers = ('Device.Time.NTPServer1', 'Device.Time.NTPServer2')
        self.get(7, prefixes=('Device.Time.',), names=ntp_servers, nonempty=ntp_servers, notes='SOLA LETTURA NTP: verifica GET di NTPServer1/NTPServer2 presenti e valorizzati. Nessun cambio server o test di sincronizzazione dopo modifica.')
        self.get(9, names=(INTERVAL, PIT, 'Device.ManagementServer.PeriodicInformEnable'), notes='GET dei parametri PI. Le prove temporali e di scrittura sono rendicontate nelle righe 8, 10 e 11.')
        for row, names in [(18, ['Device.ManagementServer.ConnectionRequestPassword']), (56, [p for p in self.parameters if p.startswith('Device.WiFi.AccessPoint.') and p.endswith('Security.KeyPassphrase')])]:
            nodes = self.matched(names=names)
            if len(nodes) != len(names) or not names:
                self.result(row, 'BLOCCATO', notes='Campi segreti non disponibili per una risposta GPV verificabile.')
                continue
            missing_response = set()
            for line in a.lines():
                if 'Missing parameter in response' in line and a.device_id in line and (parsedate(line.split()[0]) >= a.refreshed_at):
                    missing_response.update((p for p in names if p in line))
            if missing_response:
                self.result(row, 'BLOCCATO', notes='Risposta GPV incompleta: assenza del parametro non equivale a stringa vuota restituita.')
                continue
            exposed = [p for p, n in nodes.items() if n.get('_value') not in ('', None)]
            values = {p: 'NON VUOTO (valore non esportato)' if p in exposed else 'stringa vuota' for p in nodes}
            self.result(row, 'FAIL' if exposed else 'PASS', values, 'Controllo sul GPV fresco acquisito da GenieACS: campi non vuoti=' + str(len(exposed)) + '. Nessuna credenziale esportata; la censura del report non modifica la risposta CPE.')
        for row in [26, 27, 28]:
            self.result(row, 'BLOCCATO', notes='Parametri multithread/NumberOfConnections non osservati; non si sostituisce la prova con un trasferimento singolo.')
        params = self.parameters
        absent = [p for p, n in params.items() if '_value' not in n]
        stale = [p for p, n in params.items() if not n.get('_timestamp') or parsedate(n['_timestamp']) < a.refreshed_at - dt.timedelta(seconds=1)]
        self.result(31, 'PASS' if not absent and (not stale) else 'PARZIALE', {'parametri letti': len(params), 'valori mai letti': len(absent), 'timestamp obsoleti': len(stale)}, 'Elenco completo nel foglio Valori_TCL e nelle slide Data model completo. Scope: tutti i parametri restituiti dal CPE, non ogni parametro opzionale previsto dallo standard.')
        self.full_model = {p: redact_secrets(display_value(p, n), a.secrets) for p, n in params.items()}
        self.apn_values()
        self.get(33, prefixes=('Device.Cellular.Interface.1.X_CUSTOM_RadioInfo',), partial=True, notes='GET LTE/NR eseguito. Composizione GlobalCellID e formati sono da definire per TCL; non applicata la formula standard.')
        self.get(34, names=('Device.Cellular.Interface.1.X_CUSTOM_ConnectedBand',), partial=True, notes='GET bande connesse. Il significato per carrier aggregation e lo stato dei singoli component carrier richiedono conferma TCL.')
        self.get(35, prefixes=('Device.Cellular.Interface.1.X_CUSTOM_RadioInfo',), partial=True, notes='GET radio LTE/NR. Campi NR vuoti e bandwidth cellulare non esposta: dati assenti non trasformati in PASS di copertura completa.')
        self.get(38, prefixes=('Device.WiFi.AccessPoint.', 'Device.Hosts.'), partial=True, notes='GET contatori associati e Hosts. Con zero client Wi-Fi non sono verificabili i dettagli delle istanze AssociatedDevice.')
        self.get(39, prefixes=('Device.Hosts.',), notes='GET client LAN e riferimenti interfaccia. Il confronto con un inventario esterno non fa parte di questa lettura.')
        self.statistics(first_sample)
        self.get(52, names=tuple((p for p in self.parameters if p.startswith('Device.WiFi.') and p.endswith(('.Enable', '.Status', '.SupportedFrequencyBands', '.OperatingFrequencyBand', '.SSIDReference', '.LowerLayers')))), notes='GET stati, bande, SSID e sicurezza; correlazione tramite LowerLayers e SSIDReference. Credenziali oscurate.')
        channel_nodes = {p: n for p, n in self.parameters.items() if p.startswith('Device.WiFi.Radio.') and p.endswith(('.AutoChannelEnable', '.Channel', '.ChannelsInUse'))}
        active_channels = [(p, n) for p, n in channel_nodes.items() if p.endswith('.Channel') and a.value(p.rsplit('.', 1)[0] + '.AutoChannelEnable') is True]
        channel_ok = bool(active_channels) and all((type(n.get('_value')) is int and n['_value'] > 0 for p, n in active_channels))
        self.result(53, 'PASS' if channel_ok else 'FAIL' if active_channels else 'BLOCCATO', channel_nodes, 'GET in configurazione autochannel attiva: criterio TOL Channel diverso da zero. Nessun cambio canale richiesto.')
        self.get(54, names=tuple((p for p in self.parameters if p.startswith('Device.WiFi.SSID.') and p.endswith(('.Enable', '.Status', '.Name', '.LowerLayers')))), partial=True, notes='SOLA LETTURA della configurazione corrente; il requisito di DEFAULT richiede baseline di fabbrica. Nessun reset eseguito.')
        self.get(55, names=tuple((p for p in self.parameters if p.startswith('Device.WiFi.') and p.endswith(('.SSID', '.SSIDReference', '.ModeEnabled', '.KeyPassphrase')))), partial=True, notes='GET configurazione corrente main/guest e bande. Configurazione di fabbrica e uguaglianza funzionale delle password non provate tramite valori pubblici.')
        self.get(58, names=tuple((p for p in self.parameters if p.startswith('Device.WiFi.AccessPoint.') and p.endswith(('.ModeEnabled', '.ModesSupported')))), notes='GET WPA, modi dichiarati supportati e modo corrente; password oscurate.')
        modes = [(p, n) for p, n in self.parameters.items() if p.endswith('.Security.ModeEnabled')]
        if any((n.get('_value') not in [x.strip() for x in str(a.value(p.rsplit('.', 1)[0] + '.ModesSupported')).split(',')] for p, n in modes)):
            self.results[58]['status'] = 'FAIL'
            self.results[58]['notes'] += ' Modo corrente non incluso fra quelli supportati.'
        self.result(61, 'NON ESEGUITO', notes='Test CPE_0046 IMS non eseguito, come richiesto dall utente. Nessuna registrazione IMS verificata.', scope='nessuno')
        required = ['Device.DeviceInfo.' + x for x in ['Manufacturer', 'ManufacturerOUI', 'ProductClass', 'ModelName', 'SerialNumber', 'SoftwareVersion', 'HardwareVersion', 'UpTime']]
        self.get(62, prefixes=('Device.DeviceInfo.',), names=required, nonempty=required, notes='GET informazioni TCL; nomi/versioni reali. APN management dedicato non osservato.')
        self.get(63, prefixes=('Device.Cellular.Interface.1.USIM.',), notes='GET USIM; PIN oscurato. Campi SIM vuoti riportati come tali, senza inventare numeri o interpretarli come retrieve fallita.')
        self.get(64, prefixes=('Device.Ethernet.', 'Device.DHCPv4.', 'Device.IP.Interface.'), notes='GET LAN, porte e DHCP. Gli Enable Ethernet fisici RO non sono stati modificati.')
        self.get(65, prefixes=('Device.Cellular.', 'Device.IP.Interface.'), partial=True, notes='La riga Get WLAN contiene requisiti WAN/mobile. GET dei campi TCL disponibili; schema operatore e mapping completo da confermare.')
        self.get(66, prefixes=('Device.VoiceService.',), partial=True, notes='SOLA LETTURA SIP/VoIP. Configurazione credenziali, associazione porte voce e chiamate di prova non eseguite.')
        self.get(67, prefixes=('Device.DHCPv4.Server.',), notes='GET dei campi DHCP richiesti dal TOL; nessuna modifica al servizio DHCP.')
        self.get(68, prefixes=('Device.Ethernet.Interface.', 'Device.Cellular.Interface.'), partial=True, notes='Modello Device:2 TCL. Vari campi cellulari del TOL standard non sono esposti con la stessa semantica: copertura GET parziale.')
        self.get(74, names=('Device.DeviceInfo.IMEI', 'Device.Cellular.Interface.1.IMEI'), nonempty=('Device.DeviceInfo.IMEI', 'Device.Cellular.Interface.1.IMEI'), notes='GET IMEI e coerenza dei due nodi; confronto etichetta non eseguito.')
        imei1 = a.value('Device.DeviceInfo.IMEI')
        imei2 = a.value('Device.Cellular.Interface.1.IMEI')
        if imei1 != imei2 or not re.fullmatch('\\d{15}', str(imei1)):
            self.results[74]['status'] = 'FAIL'
            self.results[74]['notes'] += ' Formato/coerenza IMEI non validi.'
        self.get(75, names=tuple((p for p in self.parameters if p.startswith('Device.WiFi.') and ('.AssociatedDevice.' in p or p.endswith(('.Utilization', '.CurrentOperatingChannelBandwidth'))))), partial=True, notes='GET banda/canale/utilization/statistiche presenti; metriche per client da verificare con stazione associata e nuova discovery.')
        capability = a.node('Device.ManagementServer.InstanceWildcardsSupported')
        self.result(76, 'BLOCCATO', {'InstanceWildcardsSupported': capability.get('_value', 'non esposto')}, 'Non si inviano wildcard CWMP senza capability esplicita. L espansione dei percorsi lato ACS non dimostra il supporto sul CPE.')
        for row, note in [(14, 'VLAN non osservate'), (77, 'Enable Ethernet fisici RO'), (7, '')]:
            if row != 7:
                self.results[row]['notes'] = note + '; nessuna prova attiva equivalente inventata.'
        for row in [21, 24, 25, 29, 78]:
            self.results[row]['status'] = 'BLOCCATO'
            self.results[row]['notes'] = {21: 'Upload disponibile nel DM; manca un endpoint upload di prova validato.', 24: 'UDP echo disponibile nel DM; generazione e percorso UDP da predisporre.', 25: 'Mappatura APN-bearer-interface TR143 da validare.', 29: 'Sequenza completa UL/DL subordinata al test upload.', 78: 'AccessPoint.2.Type=Voip, Enable RW. Nessuna modifica APN: verifica funzionale VoIP fuori dal GET richiesto.'}[row]

    def statistics(self, first_sample):
        for row in range(40, 52):
            metric = ['BytesSent', 'BytesReceived', 'PacketsSent', 'PacketsReceived', 'ErrorsSent', 'ErrorsReceived'][(row - 40) % 6]
            prefix = 'Device.Cellular.Interface.' if row < 46 else 'Device.Ethernet.Interface.'
            values = {p: n for p, n in self.parameters.items() if p.startswith(prefix) and p.endswith('.Stats.' + metric)}
            current = {p: counter_value(n) for p, n in values.items()}
            previous = {p: counter_value(n) for p, n in first_sample.items() if p in values}
            bad = [p for p, v in current.items() if v is None]
            deltas = {p: v - previous[p] for p, v in current.items() if v is not None and previous.get(p) is not None}
            decreases = [p for p, v in deltas.items() if v < 0]
            note = 'Due GPV distinti; delta: ' + str(deltas) + '. Contatori zero validi. unsignedLong letto come stringa numerica ACS e interpretato senza perdita di precisione. Nessuna misura di accuratezza sotto carico.'
            status = 'BLOCCATO' if not values else 'FAIL' if bad else 'PARZIALE' if decreases else 'PASS'
            if bad:
                note += ' Contatori con tipo/valore non valido: ' + ', '.join(bad)
            if decreases:
                note += ' Reset/rollover osservato: delta da interpretare.'
            self.result(row, status, values, note)

    def apn_values(self):
        mobile_roots = [p.rsplit('.', 1)[0] + '.' for p, n in self.parameters.items() if p.startswith('Device.IP.Interface.') and p.endswith('.LowerLayers') and str(n.get('_value', '')).startswith('Device.Cellular.Interface.')]
        self.get(32, prefixes=('Device.Cellular.AccessPoint.', *mobile_roots), partial=True, notes='GET profili data/Voip e IP delle interfacce riferite al cellulare tramite LowerLayers. I riferimenti AccessPoint.Interface sono vuoti: associazione separata IP-APN VoIP non dimostrata. APN management dedicato non osservato.')

    def attributes(self):
        a = self.acs
        paths = ['Device.DeviceInfo.SoftwareVersion', 'Device.DeviceInfo.ProvisioningCode', 'Device.ManagementServer.ConnectionRequestURL']
        wan = [p for p in self.parameters if p.startswith('Device.IP.Interface.') and p.endswith('.IPv4Address.IPAddress') and (a.value(p.split('.IPv4Address')[0] + '.LowerLayers') == 'Device.Cellular.Interface.1.') and a.value(p)]
        if wan:
            paths.append(wan[0])
        try:
            a.refresh(paths, attributes=True)
            for row, path in zip([70, 71, 72, 73], paths):
                value = a.node(path).get('_notification')
                self.result(row, 'PASS' if value == 2 else 'FAIL' if value is not None else 'BLOCCATO', {path + '.Notification': value if value is not None else 'non restituito'}, 'GetParameterAttributes reale tramite provision refresh notification; atteso Notification=2.')
            if not wan:
                self.result(73, 'BLOCCATO', notes='Indirizzo del servizio cellulare non identificato.')
        except Exception as error:
            for row in [70, 71, 72, 73]:
                self.result(row, 'BLOCCATO', notes='Lettura attributi non disponibile: ' + str(error))

    def execute_active(self, row, action):
        if not self.active:
            return
        if self.acs.restore_failed:
            self.result(row, 'BLOCCATO', notes='Scritture sospese per ripristino non verificato.', scope='ATTIVO')
            return
        try:
            action()
        except TestBlocked as error:
            self.result(row, 'BLOCCATO', notes=str(error), scope='ATTIVO')
        except TimeoutError as error:
            self.result(row, 'BLOCCATO', notes='Timeout ambiente/diagnostica: ' + str(error), scope='ATTIVO')
        except Exception as error:
            self.result(row, 'FAIL', notes=type(error).__name__ + ': ' + str(error), scope='ATTIVO')

    def interval_policy(self):
        """Individua il valore costante imposto dal preset specifico di questa CPE."""
        a = self.acs
        _, presets = a.api('/presets')
        _, provisions = a.api('/provisions')
        scripts = {p['_id']: p.get('script', '') for p in provisions}
        serial = a.device.get('_deviceId', {}).get('_SerialNumber')
        for preset in presets:
            raw_condition = preset.get('precondition', '')
            try:
                condition = json.loads(raw_condition or '{}')
            except (TypeError, ValueError):
                condition = raw_condition
            serial_match = re.fullmatch('DeviceID\\.SerialNumber\\s*=\\s*"([^"]+)"', str(condition))
            applies = not condition or condition == {'DeviceID.SerialNumber': serial} or (serial_match and serial_match[1] == serial)
            if applies:
                for config in preset.get('configurations', []):
                    name = config.get('name')
                    script = scripts.get(name, '')
                    pattern = 'declare\\(\\s*["\\\']' + re.escape(INTERVAL) + '["\\\']\\s*,\\s*\\{[^}]*\\}\\s*,\\s*\\{\\s*value\\s*:\\s*(\\d+)\\s*\\}'
                    match = re.search(pattern, script)
                    if match:
                        return (int(match[1]), str(preset['_id']) + '/' + str(name))
        return (None, None)

    def periodic_600(self):
        a = self.acs
        original = a.value(INTERVAL)
        forced, policy = self.interval_policy()
        if forced is not None and forced != 60:
            note = f'Preset ACS {policy} impone {forced} s; prova TOL a 60 s per 10 minuti non eseguibile senza cambiare il preset. Nessuna modifica del preset effettuata.'
            self.result(10, 'BLOCCATO', {INTERVAL: original}, note, scope='ATTIVO')
            raise TestBlocked(note)
        if a.value('Device.ManagementServer.PeriodicInformEnable') is not True:
            raise TestBlocked('PeriodicInformEnable non attivo; non si modifica implicitamente')
        changed = False
        try:
            a.change({INTERVAL: 60})
            changed = True
            measured = a.observe(60, 10)
        finally:
            a.restore([INTERVAL])
            if changed:
                self.result(10, 'PASS', {'valore iniziale': original, 'valore temporaneo': 60, 'valore ripristinato': a.value(INTERVAL)}, 'SPV con variazione reale e ripristino verificato; cadenza rendicontata nella riga 8.', scope='ATTIVO')
        values = {'PeriodicInformInterval durante prova': 60, 'PeriodicInformInterval ripristinato': original, 'cicli misurati': len(measured['intervals_s']), 'intervalli misurati (s)': json.dumps(measured['intervals_s']), 'durata fra primo e ultimo Inform (s)': sum(measured['intervals_s'])}
        self.result(8, 'PASS' if measured['passed'] else 'FAIL', values, 'Prova TOL completa: 11 eventi 2 PERIODIC, 10 intervalli a 60 +/-5 s, retry=0. Configurazione iniziale ripristinata.', scope='ATTIVO')
        self.result(10, 'PASS', {'valore iniziale': original, 'valore temporaneo': 60, 'valore ripristinato': a.value(INTERVAL)}, 'SPV con cambio effettivo a 60 e rilettura, poi ripristino. La cadenza e verificata separatamente nella riga 8.', scope='ATTIVO')

    def periodic_time(self):
        a = self.acs
        interval = int(a.value(INTERVAL))
        new_time = (utcnow() + dt.timedelta(seconds=17)).replace(microsecond=0).isoformat().replace('+00:00', 'Z')
        initial = a.value(PIT)
        try:
            a.change({PIT: new_time})
            measured = a.observe(interval, 2)
            errors = [min((parsedate(r['timestamp']) - parsedate(new_time)).total_seconds() % interval, interval - (parsedate(r['timestamp']) - parsedate(new_time)).total_seconds() % interval) for r in measured['informs']]
        finally:
            a.restore([PIT])
        self.result(11, 'PASS' if measured['passed'] and all((x <= 5 for x in errors)) else 'FAIL', {'riferimento temporaneo': new_time, 'riferimento ripristinato': initial, 'intervalli (s)': json.dumps(measured['intervals_s']), 'scarto fase (s)': json.dumps([round(x, 3) for x in errors])}, 'PeriodicInformTime cambiato e riletto. Cadenza: ' + ('OK' if measured['passed'] else 'KO') + '; fase entro 5 s: ' + ('OK' if all((x <= 5 for x in errors)) else 'KO') + '. Scarti fase (s): ' + str([round(x, 3) for x in errors]) + '. Valore iniziale ripristinato.', scope='ATTIVO')

    def provisioning(self):
        a = self.acs
        path = 'Device.DeviceInfo.ProvisioningCode'
        description = str(self.sheet.cell(15, 6).value or '')
        variants = re.findall('(?m)^-([^\\n\\r]+)', description)
        if len(variants) != 3:
            raise TestBlocked('Le tre stringhe di prova ProvisioningCode non sono ricavabili dal TOL')
        variants = ['' if x.strip().lower() == 'null' else x.strip() for x in variants]
        tested = {}
        try:
            for value in variants:
                a.change({path: value})
                tested['stringa vuota' if value == '' else value] = 'SPV e rilettura CPE coerenti'
        finally:
            a.restore([path])
        self.result(15, 'PASS', tested, 'Tre valori del TOL provati. null interpretato come stringa vuota xsd:string, non NULL SOAP. Ripristino verificato.', scope='ATTIVO')

    def advertise(self):
        a = self.acs
        changes = {p: not n['_value'] for p, n in self.parameters.items() if p.startswith('Device.WiFi.AccessPoint.') and p.endswith('.SSIDAdvertisementEnabled') and n.get('_writable') and (type(n.get('_value')) is bool)}
        if not changes:
            raise TestBlocked('Nessun SSIDAdvertisementEnabled scrivibile e valorizzato')
        try:
            a.change(changes)
            values = {p: 'iniziale=' + str(a.protection[p][0]) + '; temporaneo=' + str(v) + '; ripristinato=' + str(a.protection[p][0]) for p, v in changes.items()}
        finally:
            a.restore(changes)
        self.result(57, 'PASS', values, 'Toggle reale SPV/GPV e ripristino. Non costituisce scansione radio dell SSID nascosto.', scope='ATTIVO')

    def wpa(self):
        a = self.acs
        changes = {}
        for p, n in self.parameters.items():
            if not p.endswith('.Security.ModeEnabled') or not n.get('_writable'):
                continue
            supported = [x.strip() for x in str(a.value(p.rsplit('.', 1)[0] + '.ModesSupported')).split(',')]
            choices = [v for v in ['WPA3-Personal', 'WPA2-WPA3-Personal', 'WPA2-Personal'] if v in supported and v != n.get('_value')]
            if choices:
                changes[p] = choices[0]
        if not changes:
            raise TestBlocked('Nessun modo alternativo WPA2/WPA3 dichiarato supportato')
        try:
            a.change(changes)
        finally:
            a.restore(changes)
        values = {p: str(v) + ' (temporaneo); ripristinato=' + str(a.value(p)) for p, v in changes.items()}
        self.result(59, 'PASS', values, 'Modi WPA2/WPA3 realmente variati, riletti e ripristinati; nessuna password modificata. Autenticazione client non provata.', scope='ATTIVO')

    def diagnostic_target(self):
        target = url.urlsplit(DOWNLOAD_URL)
        if target.scheme not in ('http','https') or not target.hostname or target.username or target.password:
            raise TestBlocked('URL TR143 HTTP/HTTPS valido senza credenziali richiesto')
        return DOWNLOAD_URL

    def diagnostic(self, row, kind):
        a = self.acs
        root = 'Device.IP.Diagnostics.' + kind
        if a.value(root + '.DiagnosticsState') == 'Requested':
            raise TestBlocked('Diagnostica gia in corso, non viene interrotta')
        target = self.diagnostic_target()
        if kind == 'DownloadDiagnostics':
            try:
                with http.urlopen(http.Request(target, method='HEAD'), timeout=10) as response:
                    size = int(response.headers.get('Content-Length', '0'))
                expected = (BASE / 'tr143_server' / 'test_50MB.bin').stat().st_size
            except (urllib.error.URLError, OSError, ValueError) as error:
                raise TestBlocked('Endpoint/file TR143 non disponibile: ' + type(error).__name__) from None
            if size != expected:
                raise TestBlocked('Dimensione HTTP diversa dal file esistente')
            changes = {root + '.DownloadURL': target}
            metrics = ['DiagnosticsState', 'TestBytesReceived', 'TotalBytesReceived', 'BOMTime', 'EOMTime', 'Interface']
        else:
            host = url.urlsplit(target).hostname
            changes = {root + '.Host': host, root + '.NumberOfRepetitions': 3, root + '.Timeout': 2000}
            metrics = ['DiagnosticsState', 'SuccessCount', 'FailureCount', 'AverageResponseTime', 'MinimumResponseTime', 'MaximumResponseTime', 'Interface']
        paths = list(changes)
        try:
            a.change({**changes, root + '.DiagnosticsState': 'Requested'})
            a.protection.pop(root + '.DiagnosticsState', None)
            deadline = time.monotonic() + 150
            while time.monotonic() < deadline:
                a.refresh([root + '.' + x for x in metrics])
                state = a.value(root + '.DiagnosticsState')
                if state not in ['Requested', 'None', '']:
                    break
                time.sleep(2)
            else:
                raise TimeoutError('Diagnostica ancora non terminale')
            values = {root + '.' + name: a.node(root + '.' + name) for name in metrics}
            success_states = ['Completed'] if kind == 'DownloadDiagnostics' else ['Complete', 'Completed']
            if state in success_states:
                if kind == 'DownloadDiagnostics':
                    bom = parsedate(a.value(root + '.BOMTime'))
                    eom = parsedate(a.value(root + '.EOMTime'))
                    seconds = (eom - bom).total_seconds()
                    good = counter_value(a.node(root + '.TestBytesReceived')) == expected and seconds > 0 and (bom.year > 2000)
                    values['dimensione attesa bytes'] = expected
                    values['durata misurata s'] = seconds
                else:
                    good = a.value(root + '.SuccessCount') == 3 and a.value(root + '.FailureCount') == 0
                status = 'PASS' if good else 'FAIL'
                if kind == 'IPPing' and a.value(root + '.SuccessCount') == 0 and (a.value(root + '.FailureCount') == 3):
                    status = 'BLOCCATO'
            elif state in ['Error_Timeout', 'Error_InitConnectionFailed', 'Error_NoResponse', 'Error_CannotResolveHostName']:
                status = 'BLOCCATO'
            else:
                status = 'FAIL'
            note = 'Destinazione ricavata dallo script di prova esistente. Stato terminale e GPV freschi. '
            if kind == 'IPPing' and status == 'BLOCCATO':
                note += 'Nessuna risposta ICMP: raggiungibilita della destinazione dall interfaccia selezionata non verificata; misura di latenza non disponibile. '
            note += 'TR143 sul payload locale: verifica diagnostica/HTTP, non misura della banda WAN.' if kind == 'DownloadDiagnostics' else 'Ping a IP di laboratorio; un timeout non e automaticamente un errore del data model.'
        finally:
            a.protection.pop(root + '.DiagnosticsState', None)
            a.restore(paths)
        self.result(row, status, values, note, scope='ATTIVO')

    def fresh_periodic_evidence(self):
        a=self.acs
        a.refresh([INTERVAL,PIT,'Device.ManagementServer.PeriodicInformEnable'])
        if a.value('Device.ManagementServer.PeriodicInformEnable') is not True:
            raise TestBlocked('PeriodicInform non abilitato')
        proof=capture_periodic_evidence(a,int(a.value(INTERVAL)),INPUT_PCAP)
        a.fetch()
        self.periodic_evidence=a.folder/'periodic'
        if not proof['passed']:
            self.results[9]['status']='FAIL'
            self.results[9]['notes']+=' PCAP non conforme: '+str(proof['intervals_s'])
        return proof

    def export(self):
        apply_manual_reviews(self)
        ws = self.sheet
        if ws.max_column > 9:
            ws.delete_cols(10, ws.max_column - 9)
        for key in list(ws.column_dimensions):
            if openpyxl.utils.column_index_from_string(key) > 11:
                del ws.column_dimensions[key]
        ws['J1'] = 'Valore parametro / esito TCL'
        ws['K1'] = 'Note TCL'
        for c in [ws['J1'], ws['K1']]:
            c._style = copy(ws['I1']._style)
            c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        ws.column_dimensions['J'].width = 110
        ws.column_dimensions['K'].width = 100
        details = self.workbook.create_sheet('Valori_TCL')
        headers = ['Riga TOL', 'TEST ID', 'Parametro / campo', 'Valore letto o misura', 'Esito', 'Ambito', 'Timestamp GPV / prova UTC']
        details.append(headers)
        for cell in details[1]:
            cell.font = Font(name='Arial', size=10, bold=True, color='FFFFFF')
            cell.fill = PatternFill('solid', fgColor='003366')
        generated = utcnow().isoformat()
        for row in self.rows:
            record = self.results[row]
            values = record['values']
            start = details.max_row + 1
            if row == 31:
                values = {**values, **self.full_model}
            for path, value in values.items():
                stamp = self.acs.node(path).get('_timestamp', generated)
                chunks = [str(value)[i:i + 29000] for i in range(0, len(str(value)), 29000)] or ['']
                for index, chunk in enumerate(chunks, 1):
                    label = path + (f' (continua {index}/{len(chunks)})' if len(chunks) > 1 else '')
                    details.append([row, self.sheet.cell(row, 2).value or 'senza ID', label, chunk, record['status'], record['scope'], stamp])
                    details.cell(details.max_row, 4).data_type = 's'
            details.append([row, self.sheet.cell(row, 2).value or 'senza ID', '[NOTE]', record['notes'], record['status'], record['scope'], generated])
            details.cell(details.max_row, 4).data_type = 's'
            text = record['status'] + '\n' + '\n'.join((f'{p}={v}' for p, v in record['values'].items()))
            if len(text) > 32700:
                text = record['status'] + '\n' + str(len(record['values'])) + ' parametri: valori integrali nel foglio Valori_TCL.'
            ws.cell(row, 10).value = text.rstrip()
            ws.cell(row, 11).value = record['notes']
            for col in [10, 11]:
                c = ws.cell(row, col)
                c.font = Font(name='Arial', size=10, color='222222')
                c.alignment = Alignment(vertical='top', wrap_text=True)
                c.data_type = 's'
                c.hyperlink = None
        for status, color in [('PASS', 'E2F0D9'), ('FAIL', 'FCE4D6'), ('BLOCCATO', 'FFF2CC'), ('PARZIALE', 'FFF2CC')]:
            ws.conditional_formatting.add('J2:J79', FormulaRule(formula=[f'LEFT(J2,{len(status)})="{status}"'], fill=PatternFill('solid', fgColor=color)))
        details.freeze_panes = 'D2'
        details.auto_filter.ref = f'A1:G{details.max_row}'
        details.sheet_view.showGridLines = False
        for key, width in [('A', 12), ('B', 25), ('C', 85), ('D', 105), ('E', 20), ('F', 15), ('G', 31)]:
            details.column_dimensions[key].width = width
        for cells in details.iter_rows(min_row=2):
            for c in cells:
                c.font = Font(name='Arial', size=10)
                c.alignment = Alignment(vertical='top', wrap_text=True)
            length = max(len(str(cells[2].value or '')) / 80, len(str(cells[3].value or '')) / 100)
            details.row_dimensions[cells[0].row].height = min(409, max(29, math.ceil(length) * 14))
        ws.auto_filter.ref = 'A1:K79'
        ws.freeze_panes = 'C2'
        ws.print_area = f'A1:K{max(self.rows)}'
        ws.print_title_rows = '1:1'
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = 'landscape'
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth = 2
        ws.page_setup.fitToHeight = 0
        identifier = getattr(self.acs, 'device_id', self.acs.device.get('_id', ''))
        self.workbook.properties.identifier = hashlib.sha256(identifier.encode()).hexdigest()
        self.workbook.save(OUTPUT)
        check = openpyxl.load_workbook(OUTPUT)
        original = openpyxl.load_workbook(SOURCE)
        for row in original['MASTER'].iter_rows(max_col=9):
            for old in row:
                new = check['MASTER'][old.coordinate]
                if old.value != new.value or old._style != new._style:
                    raise RuntimeError('Colonna originale modificata: ' + old.coordinate)
        if str(check['MASTER'].merged_cells) != str(original['MASTER'].merged_cells):
            raise RuntimeError('Unioni originali modificate')
        if hashlib.sha256(SOURCE.read_bytes()).hexdigest() != self.source_hash:
            raise RuntimeError('Il TOL sorgente e cambiato durante la prova')
        summary = {'generated_at': generated, 'source': SOURCE.name, 'source_sha256': self.source_hash, 'output': OUTPUT.name, 'rows': len(self.rows), 'source_columns_preserved': 'A:I', 'counts': dict(collections.Counter((r['status'] for r in self.results.values()))), 'active': self.active, 'restore_failed': self.acs.restore_failed, 'restorations_pending': list(self.acs.protection)}
        (self.acs.folder / 'results.json').write_text(json.dumps({'summary': summary, 'results': self.results}, ensure_ascii=False, indent=2))
        (BASE / 'logs' / 'tol_tcl_latest.json').write_text(json.dumps({'folder': str(self.acs.folder), **summary}, ensure_ascii=False, indent=2))
        self.acs.event('tol_exported', **summary)
        self.exported = True
        return summary

def run_tol(active=False):
    """Letture di default; scritture reversibili soltanto con opzione esplicita."""
    folder = BASE / 'logs' / ('tol_tcl_' + utcnow().strftime('%Y%m%dT%H%M%SZ'))
    folder.mkdir(parents=True)
    lock = (BASE / 'logs' / 'tol_tcl.lock').open('w')
    try:
        fcntl.flock(lock, fcntl.LOCK_EX | fcntl.LOCK_NB)
    except BlockingIOError:
        lock.close()
        raise TestBlocked('Un altra esecuzione TOL e gia attiva')
    a = None
    suite = None
    old_handlers = {}

    def interrupted(signum, frame):
        raise KeyboardInterrupt('Interruzione richiesta: ripristino delle modifiche TOL')
    for sig in (signal.SIGTERM, signal.SIGINT):
        old_handlers[sig] = signal.getsignal(sig)
        signal.signal(sig, interrupted)
    try:
        a = ACS(folder)
        suite = TOLSuite(a, active)
        a.event('tol_started', active=active, rows=len(suite.rows), folder=str(folder))
        a.refresh()
        a.secrets = credential_values(datamodel_parameters(a.device))
        first = {p: dict(n) for p, n in datamodel_parameters(a.device) if '.Stats.' in p}
        stats = list(first) + ['Device.DeviceInfo.UpTime']
        a.refresh(stats)
        suite.read_tests(first)
        suite.attributes()
        if active:
            suite.execute_active(8, suite.periodic_600)
            suite.execute_active(11, suite.periodic_time)
            suite.execute_active(15, suite.provisioning)
            check_paths = [p for p, n in datamodel_parameters(a.device) if p.startswith('Device.WiFi.AccessPoint.') and p.endswith('.AssociatedDeviceNumberOfEntries') or (p.startswith('Device.Ethernet.Interface.') and p.endswith('.Status'))]
            a.refresh(check_paths)
            wired = any((p.startswith('Device.Ethernet.Interface.') and p.endswith('.Status') and (str(n.get('_value')).upper() == 'UP') for p, n in datamodel_parameters(a.device)))
            clients = sum((n.get('_value', 0) for p, n in datamodel_parameters(a.device) if p.startswith('Device.WiFi.AccessPoint.') and p.endswith('.AssociatedDeviceNumberOfEntries') and (type(n.get('_value')) is int)))
            if wired and clients == 0:
                suite.execute_active(57, suite.advertise)
                suite.execute_active(59, suite.wpa)
            else:
                for row in [57, 59]:
                    suite.result(row, 'BLOCCATO', notes='Prove Wi-Fi sospese: gestione cablata non verificata o client associati presenti.', scope='ATTIVO')
            suite.execute_active(20, lambda: suite.diagnostic(20, 'DownloadDiagnostics'))
            suite.execute_active(30, lambda: suite.diagnostic(30, 'IPPing'))
        if a.protection:
            a.restore(list(a.protection))
        a.refresh()
        suite.read_tests(first)
        suite.attributes()
        try:
            measured = suite.fresh_periodic_evidence()
            a.event('final_periodic', interval_s=measured['expected_interval_s'], intervals_s=measured['intervals_s'], passed=measured['passed'])
        except Exception as error:
            suite.results[9]['notes'] += ' Osservazione temporale finale non completata: ' + str(error)
            suite.results[9]['status'] = 'PARZIALE'
        suite.export()
        return a.device
    finally:
        if a:
            if a.protection:
                try:
                    a.restore(list(a.protection))
                except Exception:
                    pass
            a.cleanup()
            if suite:
                try:
                    if not suite.exported or a.restore_failed:
                        suite.export()
                except Exception as error:
                    a.event('export_incomplete', reason=type(error).__name__)
        fcntl.flock(lock, fcntl.LOCK_UN)
        lock.close()
        for sig, handler in old_handlers.items():
            signal.signal(sig, handler)
# tr069_validator_agent
LOG_FILE = PERIODIC_LOG_PATH
EXCEL_FILE = EXCEL_PATH

def get_param_value(dev, path):
    parts = path.split('.')
    curr = dev
    for p in parts:
        if isinstance(curr, dict) and p in curr:
            curr = curr[p]
        else:
            return None
    if isinstance(curr, dict) and '_value' in curr:
        return curr['_value']
    return curr

def run_test_suite(device=None):
    print('[*] Estrazione dati CPE dall ACS...')
    dev = device if device is not None else get_mongo_device()
    if not dev:
        print('[!] Errore: Impossibile estrarre i dati da MongoDB.')
        return
    secrets = credential_values(datamodel_parameters(dev))
    dev_id = redact_secrets(dev.get('_id', 'Unknown'), secrets)
    print(f'[+] Dispositivo rilevato: {dev_id}\n')
    ensure_audit_workbook()
    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb['Sheet2']
    interval_val = get_param_value(dev, 'Device.ManagementServer.PeriodicInformInterval') or 0
    for row in range(2, ws.max_row + 1):
        param = ws.cell(row=row, column=1).value
        expected = ws.cell(row=row, column=2).value
        if not param:
            continue
        val = get_param_value(dev, param)
        esito = 'FAIL'
        note = ''
        if param == 'Device.ManagementServer.PeriodicInformTime':
            periodic_ok, periodic_detail, _ = audit_periodic_inform_evidence(dev, int(interval_val))
            if val is None or str(val).strip() == '':
                esito, note = ('FAIL', 'FAIL: PeriodicInformTime non valorizzato')
            else:
                esito = 'PASS' if periodic_ok else 'WARN' if periodic_detail.startswith('WARN:') else 'FAIL'
                note = periodic_detail
            ws.cell(row=row, column=5).value = LOG_FILE
        elif 'DeviceID.' in param:
            did_key = param.split('.')[1]
            did_obj = dev.get('_deviceId', {})
            dval = did_obj.get(f'_{did_key}') or did_obj.get(did_key)
            if dval and (expected is None or str(expected).lower() in str(dval).lower() or 'not null' in str(expected).lower()):
                esito = 'PASS'
                note = f'OK: Corrisponde ({dval})'
            else:
                esito = 'FAIL'
                note = f'FAIL: Trovato ({dval})'
        elif 'IMEI' in param or 'SerialNumber' in param:
            if val and len(str(val).strip()) >= 10:
                esito = 'PASS'
                note = f'OK: Conforme ({val})'
            else:
                esito = 'FAIL'
                note = f'FAIL: {val}'
        elif expected is not None and str(expected).strip().upper() != 'SET':
            if str(expected).strip().lower() == 'not null':
                matches = val is not None and str(val).strip() != ''
            else:
                matches = val is not None and str(val) == str(expected)
            esito = 'PASS' if matches else 'FAIL'
            note = f'OK: {val}' if matches else f'FAIL: atteso {expected}; letto {val}'
        elif val is not None:
            esito = 'PASS'
            note = f'OK: {val}'
        else:
            esito = 'FAIL'
            note = 'FAIL: Parametro non valorizzato'
        note = redact_secrets(note, secrets)
        if SENSITIVE_PATH.search(str(param)):
            note = f'{esito}: Valore oscurato'
        ws.cell(row=row, column=3).value = note
        ws.cell(row=row, column=4).value = esito
        status_tag = f'[{esito}]'.ljust(7)
        print(f'  {status_tag} {param:<45} -> {note}')
    wb.save(EXCEL_FILE)
    print(f'\n[+] Data Model dedotto: TR-181')
    print(f'[+] Test completati. File aggiornato: {EXCEL_FILE}')
    print(f'[+] Evidenza Periodic Inform salvata in: {LOG_FILE}')
    return dev
DOWNLOAD_URL = os.environ.get('TR143_DOWNLOAD_URL', 'http://192.0.2.1:8080/test_50MB.bin')
AUDIT_BASELINE = [['Device.DeviceInfo.HardwareVersion', 'HH516L-V3.0'], ['Device.DeviceInfo.IMEI', 'not null'], ['Device.DeviceInfo.ModemFirmwareVersion', 'HH516L_FT_03.00_17'], ['Device.DeviceInfo.SoftwareVersion', 'HH516L_FT_03.00_17'], ['Device.DeviceInfo.SpecVersion', '1.0'], ['Device.ManagementServer.ConnectionRequestPassword', 'not null'], ['Device.ManagementServer.ConnectionRequestURL', 'not null'], ['Device.ManagementServer.ConnectionRequestUsername', 'not null'], ['Device.ManagementServer.PeriodicInformEnable', True], ['Device.ManagementServer.PeriodicInformInterval', 30], ['Device.ManagementServer.PeriodicInformTime', "verificare che rispetti l'invio ogni quanto definito in Device.ManagementServer.PeriodicInformInterval"], ['DeviceID.Manufacturer', 'TCL Communication'], ['DeviceID.OUI', 'not null'], ['DeviceID.ProductClass', 'TCL 5G CPE'], ['DeviceID.SerialNumber', 'not null'], ['Device.ManagementServer.PeriodicInformInterval', 'SET']]
# Acquisizione e verifica PCAP: nessun fallback che consideri tutti gli Inform periodici.
def _command(args, timeout=20):
    return subprocess.run(args, capture_output=True, text=True, timeout=timeout, check=True)


class PeriodicCapture:
    """Processo temporaneo isolato, rete host; non avvia/ferma laboratori Kathara."""
    def __init__(self, acs, folder, seconds, port):
        self.acs, self.folder, self.seconds, self.port = acs, folder, seconds, port
        self.container = None
        self.path = folder / 'capture_full.pcap'

    def start(self):
        if not shutil.which('tshark'):
            raise TestBlocked('tshark non disponibile per analizzare il PCAP')
        rows = json.loads(_command(['docker','inspect',self.acs.acs]).stdout)
        lab_prefix = rows[0]['Name'].split('_acs_')[0]
        candidates = _command(['docker','ps','-q']).stdout.split()
        image = None
        for cid in candidates:
            row = json.loads(_command(['docker','inspect',cid]).stdout)[0]
            if not row['Name'].startswith(lab_prefix+'_cpe_net_'):
                continue
            result = subprocess.run(['docker','exec',cid,'sh','-c','command -v tcpdump; command -v timeout'],capture_output=True,text=True,timeout=10)
            if result.returncode == 0 and '/usr/bin/tcpdump' in result.stdout and '/usr/bin/timeout' in result.stdout:
                image = row['Image'];break
        if not image:
            raise TestBlocked('Immagine locale del laboratorio con tcpdump/timeout non individuata')
        name = 'tr069-pcap-'+uuid.uuid4().hex[:12]
        args = ['docker','run','-d','--name',name,'--label','tr069.temporary=periodic-capture',
                '--network','host','--cap-drop','ALL','--cap-add','NET_RAW','--cap-add','SETUID','--cap-add','SETGID',
                '--entrypoint','/usr/bin/timeout',image,'--signal=INT','--kill-after=5',str(self.seconds),
                '/usr/bin/tcpdump','-Z','root','-p','-i','any','-nn','-s','0','-U','-w','/tmp/periodic.pcap',
                'tcp','port',str(self.port)]
        self.container = _command(args).stdout.strip()
        try:
            end = time.monotonic()+10
            while time.monotonic()<end:
                state = json.loads(_command(['docker','inspect',self.container]).stdout)[0]['State']
                logs = _command(['docker','logs',self.container])
                message = logs.stdout+logs.stderr
                if not state['Running']:
                    raise TestBlocked('Avvio tcpdump non riuscito: '+sanitise_text(message[-400:]))
                if 'listening on' in message:
                    self.started = utcnow()
                    self.acs.event('pcap_capture_started',helper=name,port=self.port,limit_s=self.seconds)
                    return
                time.sleep(0.2)
            raise TestBlocked('Tcpdump non conferma l avvio entro 10 s')
        except BaseException:
            self.close(copy_capture=False)
            raise

    def close(self, copy_capture=True):
        if not self.container:
            return
        cid, self.container = self.container, None
        try:
            state = json.loads(_command(['docker','inspect',cid]).stdout)[0]['State']
            if state['Running']:
                subprocess.run(['docker','kill','--signal=INT',cid],capture_output=True,text=True,timeout=10)
                try:_command(['docker','wait',cid],timeout=10)
                except subprocess.TimeoutExpired:_command(['docker','stop','--time','5',cid],timeout=15)
            logs = _command(['docker','logs',cid])
            (self.folder/'capture.log').write_text(sanitise_text(logs.stdout+logs.stderr))
            if copy_capture:
                _command(['docker','cp',cid+':/tmp/periodic.pcap',str(self.path)])
                self.path.chmod(0o600)
        finally:
            # Rimuove esclusivamente il container creato da questa istanza.
            subprocess.run(['docker','stop','--time','5',cid],capture_output=True,timeout=15)
            _command(['docker','rm',cid])


def _xml_text(root, name):
    element = root.find('.//{*}'+name)
    return element.text.strip() if element is not None and element.text else ''


def decode_pcap_messages(path, port):
    args=['tshark','-2','-r',str(path),'-d',f'tcp.port=={int(port)},http',
          '-o','tcp.desegment_tcp_streams:TRUE','-o','http.desegment_body:TRUE',
          '-Y','http.file_data','-T','fields']
    for field in ['frame.number','frame.time_epoch','tcp.stream','http.request.method','http.response.code','http.file_data']:
        args += ['-e',field]
    args += ['-E','separator=|']
    output=_command(args,timeout=40).stdout
    messages=[]
    for line in output.splitlines():
        fields=line.split('|')
        if len(fields)!=6 or not fields[5] or len(fields[5])>4000000:
            continue
        try:
            payload=bytes.fromhex(fields[5].replace(':',''))
            root=ET.fromstring(payload)
            envelope_id=_xml_text(root,'ID')
            base={'frame':int(fields[0]),'epoch':float(fields[1]),'stream':int(fields[2]),
                  'cwmp_id':envelope_id,'body_sha256':hashlib.sha256(payload).hexdigest()}
            if root.find('.//{*}Inform') is not None:
                events=[(e.text or '').strip() for e in root.findall('.//{*}EventCode')]
                retry=_xml_text(root,'RetryCount')
                messages.append({**base,'kind':'Inform','serial':_xml_text(root,'SerialNumber'),
                                 'event_codes':', '.join(events),'retry_count':int(retry) if retry else None})
            elif root.find('.//{*}InformResponse') is not None and fields[4]=='200':
                messages.append({**base,'kind':'InformResponse'})
        except (ValueError,ET.ParseError):
            continue
    return sorted(messages,key=lambda item:(item['epoch'],item['frame']))


def pcap_periodic_rows(path, port, serial):
    messages=decode_pcap_messages(path,port)
    informs=[]
    for row in messages:
        if row['kind']!='Inform' or row['serial']!=serial or '2 PERIODIC' not in [s.strip() for s in row['event_codes'].split(',')]:
            continue
        # Capture "any" vede anche il medesimo pacchetto prima/dopo il NAT Docker.
        if any(row['body_sha256']==old['body_sha256'] and abs(row['epoch']-old['epoch'])<2 for old in informs[-4:]):
            continue
        responses=[r for r in messages if r['kind']=='InformResponse' and r['stream']==row['stream']
                   and r['cwmp_id']==row['cwmp_id'] and 0<=r['epoch']-row['epoch']<=5]
        response=responses[0] if responses else None
        informs.append({k:v for k,v in row.items() if k not in {'serial','kind'}} |
                       {'timestamp':dt.datetime.fromtimestamp(row['epoch'],dt.timezone.utc).isoformat(),
                        'response_frame':response['frame'] if response else None,
                        'response_http':200 if response else None})
    return informs


def _pcap_measure(rows, interval):
    deltas=[round(b['epoch']-a['epoch'],6) for a,b in zip(rows,rows[1:])]
    passed=(len(rows)>=3 and all(0<d and abs(d-interval)<=5 for d in deltas)
            and all(r['retry_count']==0 and r['response_http']==200 for r in rows))
    return {'informs':rows,'intervals_s':deltas,'passed':passed,'expected_interval_s':interval,'tolerance_s':5}


def cwmp_port(device):
    target=url.urlsplit(str(get_param_value(device,'Device.ManagementServer.URL') or ''))
    if target.scheme!='http' or not target.hostname:
        raise TestBlocked('PCAP SOAP leggibile richiede endpoint CWMP HTTP; HTTPS richiede evidenza TLS dedicata')
    return target.port or 80


def capture_periodic_evidence(a, interval, input_pcap=None):
    if type(interval) is not int or not 1<=interval<=300:
        raise TestBlocked('Intervallo PI fuori limite 1..300 s: nessuna attesa automatica lunga')
    folder=a.folder/'periodic';folder.mkdir(mode=0o700,exist_ok=True)
    port=cwmp_port(a.device)
    serial=a.device.get('_deviceId',{}).get('_SerialNumber')
    if not serial:raise TestBlocked('Identificativo CPE non disponibile per selezionare il PCAP')
    started=utcnow();acs_rows=[];capture=None
    if input_pcap:
        full=Path(input_pcap).resolve()
        if not full.is_file():raise TestBlocked('PCAP sorgente non presente')
        rows=pcap_periodic_rows(full,port,serial)[-3:]
        source_kind='pcap'
    else:
        capture=PeriodicCapture(a,folder,min(1000,interval*5+60),port)
        try:
            capture.start()
            observed=a.observe(interval,2,since=capture.started)
            acs_rows=observed['informs']
            time.sleep(1)  # Consente a InformResponse di essere incluso nella cattura.
        finally:
            capture.close()
        full=capture.path
        available=pcap_periodic_rows(full,port,serial)
        rows=[]
        for event in acs_rows:
            stamp=parsedate(event['timestamp']).timestamp()
            matches=[r for r in available if abs(r['epoch']-stamp)<2 and r['cwmp_id']==event.get('request_id')]
            if not matches:raise TestBlocked('Inform ACS senza corrispondenza nel PCAP')
            rows.append(min(matches,key=lambda r:abs(r['epoch']-stamp)))
        source_kind='pcap+acs_log'
    if len(rows)<3:raise TestBlocked('PCAP con meno di tre Inform periodici del CPE selezionato')
    streams=sorted({int(r['stream']) for r in rows})
    target=folder/'periodic_inform.pcap'
    _command(['tshark','-r',str(full),'-Y','tcp.stream in {'+', '.join(map(str,streams))+'}',
              '-F','pcap','-w',str(target)],timeout=40)
    target.chmod(0o600)
    final_rows=pcap_periodic_rows(target,port,serial)
    if len(final_rows)!=3:raise TestBlocked('Estratto PCAP non contiene esattamente tre Inform periodici')
    measured=_pcap_measure(final_rows,interval)
    proof={'evidence_type':source_kind,**measured,'device_sha256':hashlib.sha256(a.device_id.encode()).hexdigest(),
           'started':started.isoformat(),'ended':utcnow().isoformat(),'final_interval_s':int(a.value(INTERVAL)),
           'periodic_inform_time':a.value(PIT),'periodic_inform_enable':a.value('Device.ManagementServer.PeriodicInformEnable'),
           'pcap':str(target.relative_to(BASE)),'pcap_sha256':hashlib.sha256(target.read_bytes()).hexdigest(),
           'cwmp_port':port,'acs_informs':acs_rows,'outcome':'PASS' if measured['passed'] else 'FAIL'}
    (folder/'summary.json').write_text(json.dumps(proof,indent=2))
    (BASE/'logs'/'periodic_latest.json').write_text(json.dumps({'summary':str((folder/'summary.json').relative_to(BASE))},indent=2))
    if capture and full.exists():full.unlink()
    a.event('pcap_verified',informs=len(final_rows),intervals_s=measured['intervals_s'],passed=measured['passed'],pcap=proof['pcap'])
    return proof


def verified_periodic_evidence(device, expected_interval, max_age_seconds=3600):
    candidates=list((BASE/'logs').glob('*/periodic/summary.json'))
    if not candidates:raise TestBlocked('Nessuna evidenza PCAP: eseguire una nuova acquisizione')
    source=max(candidates,key=lambda p:p.stat().st_mtime)
    proof=json.loads(source.read_text())
    if proof.get('evidence_type') not in ('pcap','pcap+acs_log'):raise TestBlocked('Fonte non PCAP')
    if proof.get('device_sha256')!=hashlib.sha256(str(device.get('_id','')).encode()).hexdigest():raise TestBlocked('Evidenza relativa a un altro CPE')
    age=(utcnow()-parsedate(proof['ended'])).total_seconds()
    if not -5<=age<=max_age_seconds:raise TestBlocked('PCAP non recente; ripetere l acquisizione')
    interval=get_param_value(device,INTERVAL)
    if proof.get('expected_interval_s')!=expected_interval or proof.get('final_interval_s')!=interval or interval!=expected_interval:
        raise TestBlocked('PCAP non coerente con PeriodicInformInterval corrente')
    if get_param_value(device,'Device.ManagementServer.PeriodicInformEnable') is not True or proof.get('periodic_inform_enable') is not True:
        raise TestBlocked('PeriodicInform non abilitato')
    if proof.get('periodic_inform_time')!=get_param_value(device,PIT):raise TestBlocked('PeriodicInformTime modificato dopo la prova PCAP')
    path=(BASE/proof['pcap']).resolve()
    if not path.is_relative_to((BASE/'logs').resolve()) or not path.is_file():raise TestBlocked('PCAP mancante o esterno alla cartella evidenze')
    if hashlib.sha256(path.read_bytes()).hexdigest()!=proof['pcap_sha256']:raise TestBlocked('Hash PCAP diverso da quello registrato')
    rows=pcap_periodic_rows(path,int(proof['cwmp_port']),device.get('_deviceId',{}).get('_SerialNumber'))
    measured=_pcap_measure(rows,expected_interval)
    if rows and not -5<=(utcnow().timestamp()-rows[-1]['epoch'])<=max_age_seconds:raise TestBlocked('Timestamp dei pacchetti PCAP non recente')
    if len(rows)!=3:raise TestBlocked('PCAP senza tre Inform distinti')
    if proof['evidence_type']=='pcap+acs_log':
        if len(proof.get('acs_informs',[]))!=3:raise TestBlocked('Log ACS incompleto')
        for row,event in zip(rows,proof['acs_informs']):
            if abs(row['epoch']-parsedate(event['timestamp']).timestamp())>=2 or row['cwmp_id']!=event.get('request_id'):
                raise TestBlocked('PCAP non correlato ai tre Inform ACS')
    return {**proof,**measured,'source':str(source.relative_to(BASE))}


def audit_periodic_inform_evidence(device, expected_interval, max_age_seconds=3600):
    try:proof=verified_periodic_evidence(device,expected_interval,max_age_seconds)
    except (TestBlocked,ValueError,KeyError,TypeError,OSError,subprocess.SubprocessError) as error:
        text='WARN: '+str(error)
        Path(LOG_FILE).write_text('AUDIT PERIODIC INFORM - PCAP\n'+text+'\n')
        return False,text,None
    status='PASS' if proof['passed'] else 'FAIL'
    lines=['AUDIT PERIODIC INFORM - PCAP E RISPOSTE ACS',
           'Fonte: '+proof['pcap'],'SHA256: '+proof['pcap_sha256'],
           f'Intervallo atteso: {expected_interval} s; tolleranza +/-5 s',
           'Tre Inform 2 PERIODIC con rispettiva InformResponse HTTP 200; retry=0 richiesto.']
    for i,row in enumerate(proof['informs']):
        delta='riferimento' if i==0 else str(proof['intervals_s'][i-1])+' s'
        lines.append(f"{i+1}: {row['timestamp']} frame {row['frame']}/{row['response_frame']}, delta {delta}, retry {row['retry_count']}")
    lines.append('VERDETTO: '+status)
    Path(LOG_FILE).write_text('\n'.join(lines)+'\n')
    return proof['passed'],status+': cicli '+' / '.join(f'{d:.3f} s' for d in proof['intervals_s']),sum(proof['intervals_s'])/len(proof['intervals_s'])


def add_pcap_evidence_slide(prs,device):
    slide=prs.slides.add_slide(prs.slide_layouts[6]);apply_background(slide,prs)
    add_header(slide,'Periodic Inform: evidenza PCAP','Verifica sul traffico CWMP del CPE; due cicli completi')
    try:proof=verified_periodic_evidence(device,int(get_param_value(device,INTERVAL) or 0))
    except (TestBlocked,ValueError,KeyError,TypeError,OSError,subprocess.SubprocessError) as error:
        box=slide.shapes.add_textbox(Inches(.8),Inches(1.8),Inches(11.73),Inches(2))
        box.text_frame.text='WARN: '+str(error)+'\nEseguire lo script senza --report-only per acquisire una prova aggiornata.'
        for p in box.text_frame.paragraphs:p.font.size=Pt(18);p.font.color.rgb=C_WARN
        return
    table=slide.shapes.add_table(4,6,Inches(.8),Inches(1.8),Inches(11.73),Inches(2.1)).table
    for c,w in zip(table.columns,[.65,3.3,1.3,1.1,1.4,3.98]):c.width=Inches(w)
    data=[['#','Timestamp UTC','Frame req/risp','TCP stream','Delta (s)','Evento / risposta']]
    for i,row in enumerate(proof['informs']):
        data.append([str(i+1),row['timestamp'],f"{row['frame']} / {row['response_frame']}",str(row['stream']),
                     '-' if i==0 else f"{proof['intervals_s'][i-1]:.6f}",f"{row['event_codes']}\nHTTP {row['response_http']}; retry {row['retry_count']}"])
    for i,row in enumerate(data):
        for j,value in enumerate(row):set_table_cell(table.cell(i,j),value,header=i==0,font_size=11)
    box=slide.shapes.add_textbox(Inches(.8),Inches(4.2),Inches(11.73),Inches(2.5));tf=box.text_frame;tf.word_wrap=True
    entries=[('PASS' if proof['passed'] else 'FAIL')+f": intervallo atteso {proof['expected_interval_s']} s; tolleranza +/-5 s.",
             'PCAP: '+proof['pcap'],'SHA256: '+proof['pcap_sha256'],
             'Cattura: '+proof['started']+' - '+proof['ended'],
             'Sono conservate soltanto le sessioni dei tre Inform. Numeri frame riferiti al PCAP estratto.']
    for i,text in enumerate(entries):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=text;p.font.name='Arial';p.font.size=Pt(14 if i==0 else 11)
        p.font.color.rgb=(C_PASS if proof['passed'] else C_FAIL) if i==0 else C_DARK;p.space_after=Pt(10)

# Ingressi e output del flusso unico.
INPUT_PCAP = None


def ensure_audit_workbook():
    if Path(EXCEL_FILE).exists():return
    wb=openpyxl.Workbook();ws=wb.active;ws.title='Sheet2'
    ws.append(['Parametro','Valore atteso / operazione','Esito / valore','Status','Evidenza / dettagli'])
    for path,expected in AUDIT_BASELINE:ws.append([path,expected])
    ws.freeze_panes='A2'
    for c in ws[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='102542')
    for col,width in [('A',65),('B',45),('C',80),('D',16),('E',65)]:ws.column_dimensions[col].width=width
    wb.save(EXCEL_FILE)


def fetch_device():
    request=http.Request(NBI+'/devices')
    with http.urlopen(request,timeout=15) as response:rows=json.load(response)
    if len(rows)!=1:raise TestBlocked('Selezione CPE non univoca')
    return rows[0]


def get_mongo_device():
    """Compatibilita interna: legge dall API NBI, senza shell o client Mongo esterni."""
    return fetch_device()


def export_datamodel(device):
    parameters=datamodel_parameters(device);secrets=credential_values(parameters)
    safe=json.loads(json.dumps(device,default=str))
    def mask(node,path=''):
        if not isinstance(node,dict):return
        if '_value' in node:
            if SENSITIVE_PATH.search(path):node['_value']='[oscurato]'
            elif isinstance(node['_value'],str):node['_value']=redact_secrets(display_value(path,node) if '://' in node['_value'] else node['_value'],secrets)
        for key,value in node.items():
            if isinstance(value,dict):mask(value,(path+'.'+key).strip('.'))
            elif isinstance(value,str):node[key]=redact_secrets(value,secrets)
            elif isinstance(value,list):node[key]=[redact_secrets(v,secrets) if isinstance(v,str) else v for v in value]
            elif isinstance(value,(int,float)) and any(secret in str(value) for secret in secrets):node[key]='[oscurato]'
    mask(safe)
    target=BASE/'cpe_datamodel_full.json';target.write_text(json.dumps(safe,ensure_ascii=False,indent=2));target.chmod(0o600)
    wb=openpyxl.Workbook();ws=wb.active;ws.title='DataModel'
    ws.append(['Parametro','Valore','Tipo XML','Accesso','Stato valore','Timestamp ACS UTC'])
    for path,node in parameters:
        ws.append([path,redact_secrets(display_value(path,node),secrets),node.get('_type',''),
                   'RW' if node.get('_writable') else 'RO',value_state(node),node.get('_timestamp','')])
        ws.cell(ws.max_row,2).data_type='s'
    for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='102542')
    for col,width in [('A',88),('B',105),('C',24),('D',12),('E',20),('F',32)]:ws.column_dimensions[col].width=width
    ws.freeze_panes='B2';ws.auto_filter.ref=ws.dimensions
    for row in ws.iter_rows(min_row=2):
        for c in row:c.alignment=Alignment(wrap_text=True,vertical='top')
        ws.row_dimensions[row[0].row].height=min(409,max(28,14*math.ceil(max(len(str(row[0].value))/88,len(str(row[1].value))/100))))
    wb.save(BASE/'cpe_datamodel_flat.xlsx')
    print(f'[+] Data model esportato: {len(parameters)} parametri, JSON e XLSX con credenziali oscurate.',flush=True)


def previous_manual_reviews(device_id, source_hash):
    """Riporta solo validazioni esplicite dell utente, vincolate a CPE e modello TOL."""
    if not OUTPUT.is_file():return {}
    try:
        meta=json.loads((BASE/'logs'/'tol_tcl_latest.json').read_text())
        if meta.get('source_sha256')!=source_hash:return {}
        wb=openpyxl.load_workbook(OUTPUT,data_only=True)
        if wb.properties.identifier!=hashlib.sha256(device_id.encode()).hexdigest():return {}
        details=wb['Valori_TCL'];reviews={}
        previous=json.loads((Path(meta['folder'])/'results.json').read_text()).get('results',{})
        for cells in details.iter_rows(min_row=2,values_only=True):
            row,test_id,parameter,value,status,scope,*_=cells
            if scope=='VALIDAZIONE UTENTE' and parameter=='[NOTE]':
                reviews[int(row)]={'test_id':test_id,'status':status,
                                   'reviewed_at':previous.get(str(row),{}).get('reviewed_at',meta.get('generated_at'))}
        return reviews
    except (OSError,ValueError,KeyError,TypeError):return {}


def apply_manual_reviews(suite):
    for row,review in getattr(suite,'manual_reviews',{}).items():
        if suite.sheet.cell(row,2).value!=review['test_id'] and not (suite.sheet.cell(row,2).value is None and str(review['test_id']).lower().replace(' ','')=='senzaid'):
            continue
        record=suite.results[row]
        if record.get('scope')=='VALIDAZIONE UTENTE':continue
        automatic=record['status']
        record['automatic_status']=automatic;record['reviewed_at']=review['reviewed_at']
        record['status']=review['status'];record['scope']='VALIDAZIONE UTENTE'
        record['notes']=(f"Esito {review['status']} validato dall utente il {review['reviewed_at']}. "
                         f"Esito automatico di questa esecuzione: {automatic}. "+record['notes'])
        if row==11 and not record['values']:
            record['values']={p:redact_secrets(display_value(p,suite.acs.node(p)),suite.acs.secrets) for p in (PIT,INTERVAL)}


def run_periodic_only(input_pcap=None):
    folder=BASE/'logs'/('run_'+utcnow().strftime('%Y%m%dT%H%M%SZ'));folder.mkdir(mode=0o700,parents=True)
    lock=(BASE/'logs'/'tol_tcl.lock').open('w');fcntl.flock(lock,fcntl.LOCK_EX|fcntl.LOCK_NB)
    a=None
    try:
        a=ACS(folder);a.refresh([INTERVAL,PIT,'Device.ManagementServer.PeriodicInformEnable'])
        proof=capture_periodic_evidence(a,int(a.value(INTERVAL)),input_pcap)
        a.fetch();export_datamodel(a.device);run_test_suite(a.device);build_presentation(a.device)
        return 0 if proof['passed'] else 2
    finally:
        if a:a.cleanup()
        fcntl.flock(lock,fcntl.LOCK_UN);lock.close()


def main(argv=None):
    import argparse
    parser=argparse.ArgumentParser(description='TOL TCL, TR-069/TR-181, PCAP Periodic Inform e PPTX in un solo script')
    mode=parser.add_mutually_exclusive_group()
    mode.add_argument('--report-only',action='store_true',help='Rigenera XLSX data model e PPTX dalla cache ACS; non esegue test nuovi')
    mode.add_argument('--periodic-only',action='store_true',help='Rilegge PI, acquisisce due cicli PCAP e rigenera il PPTX senza modificare la TOL')
    mode.add_argument('--self-test',action='store_true',help='Esegue i controlli interni isolati, senza ACS o CPE')
    parser.add_argument('--tol-active',action='store_true',help='Aggiunge le prove reversibili con rilettura e ripristino; nessun reset/firmware/APN/password')
    parser.add_argument('--pcap',type=Path,help='Usa una cattura esistente: estrae tre Inform del CPE con risposte e verifica data/intervallo')
    parser.add_argument('--template',type=Path,help='Modello PPTX di 13 slide, con le slide editoriali già preparate')
    parser.add_argument('--download-url',default=DOWNLOAD_URL,help='Endpoint TR143 del payload da 50 MiB, configurato dal testbed')
    args=parser.parse_args(argv)
    if args.tol_active and (args.report_only or args.periodic_only or args.self_test):parser.error('--tol-active richiede la suite TOL completa')
    global INPUT_PCAP, REPORT_TEMPLATE
    if args.template:REPORT_TEMPLATE=args.template.resolve()
    INPUT_PCAP=args.pcap
    globals()['DOWNLOAD_URL']=args.download_url
    (BASE/'logs').mkdir(mode=0o700,exist_ok=True)
    if args.self_test:return run_self_tests()
    def terminate(signum,frame):raise KeyboardInterrupt('Arresto richiesto: pulizia risorse temporanee')
    old_term_handler=signal.signal(signal.SIGTERM,terminate)
    try:
        if args.periodic_only:return run_periodic_only(args.pcap)
        if args.report_only:
            device=fetch_device()
        else:
            device=run_tol(active=args.tol_active)
        export_datamodel(device)
        run_test_suite(device)
        build_presentation(device)
        try:proof=verified_periodic_evidence(device,int(get_param_value(device,INTERVAL) or 0))
        except (TestBlocked,OSError,ValueError,KeyError,TypeError,subprocess.SubprocessError) as error:
            print('[!] Evidenza Periodic Inform non valida: '+str(error),flush=True)
            return 3
        return 0 if proof['passed'] else 2
    except (TestBlocked,TimeoutError,urllib.error.URLError,subprocess.SubprocessError,OSError) as error:
        print('[!] Esecuzione bloccata: '+type(error).__name__+': '+str(error),flush=True)
        return 3
    except KeyboardInterrupt:
        print('[!] Interruzione richiesta; ripristini e pulizia delle sole risorse temporanee eseguiti.',flush=True)
        return 130
    finally:
        signal.signal(signal.SIGTERM,old_term_handler)

def run_self_tests():
    """Controlli isolati: non invia task ACS e non modifica configurazioni CPE."""
    import contextlib
    import io
    import sys
    import unittest
    from types import SimpleNamespace
    from unittest.mock import patch
    tol=sys.modules[__name__]

    class FakeACS:

        def __init__(self, folder):
            self.folder = folder
            self.secrets = []
            self.device = {'Device': {'Test': {'_object': False, '_value': 0}}}
            self.protection = {}
            self.restore_failed = False

        def event(self, *args, **kwargs):
            pass

        def node(self, path):
            return {}

    class TOLTests(unittest.TestCase):

        def suite(self, folder):
            s = tol.TOLSuite(FakeACS(folder))
            s.full_model = {'Device.Test': '0'}
            for row in s.rows:
                s.results[row] = {'status': 'NON ESEGUITO', 'values': {}, 'notes': 'Non eseguito nel test isolato.', 'scope': 'nessuno'}
            s.results[9] = {'status': 'PASS', 'values': {'Device.ManagementServer.PeriodicInformInterval': '30'}, 'notes': 'GET di prova', 'scope': 'GET'}
            return s

        def test_original_columns_and_styles_preserved(self):
            with tempfile.TemporaryDirectory(dir=BASE/'logs') as tmp:
                folder = Path(tmp)
                s = self.suite(folder)
                with patch.object(tol, 'OUTPUT', folder / 'tol.xlsx'), patch.object(tol, 'BASE', folder):
                    (folder / 'logs').mkdir()
                    summary = s.export()
                    book = tol.openpyxl.load_workbook(folder / 'tol.xlsx')
                    self.assertEqual(book['MASTER'].max_column, 11)
                    self.assertEqual(book['MASTER']['J1'].value, 'Valore parametro / esito TCL')
                    self.assertEqual(summary['source_columns_preserved'], 'A:I')
                    self.assertEqual(book['MASTER']['J9'].value, 'PASS\nDevice.ManagementServer.PeriodicInformInterval=30')

        def test_terminal_diagnostic_not_rejected_as_failed_set(self):
            a = tol.ACS.__new__(tol.ACS)
            a.restore_failed = False
            a.protection = {}
            a.node = lambda p: {'_writable': True, '_type': 'xsd:string', '_value': 'Completed', '_timestamp': dt.datetime.now(dt.timezone.utc).isoformat()}
            a.post = lambda *args, **kwargs: 'test'
            a.wait = lambda *args: None
            a.event = lambda *args, **kwargs: None
            a.change({'Device.IP.Diagnostics.IPPing.DiagnosticsState': 'Requested'})

        def test_datetime_equivalence_and_sequential_read(self):
            a = tol.ACS.__new__(tol.ACS)
            a.restore_failed = False
            a.protection = {}
            events = []
            a.node = lambda p: {'_writable': True, '_type': 'xsd:dateTime', '_value': '2026-10-17T00:00:00.000Z', '_timestamp': dt.datetime.now(dt.timezone.utc).isoformat()}

            def post(task, **kwargs):
                events.append(task['name'])
                return str(len(events))
            a.post = post
            a.wait = lambda ids: events.append('wait')
            a.event = lambda *args, **kwargs: None
            a.change({tol.PIT: '2026-10-17T00:00:00Z'})
            self.assertEqual(events, ['setParameterValues', 'wait', 'getParameterValues', 'wait'])

        def test_forced_interval_policy_is_blocked(self):
            with tempfile.TemporaryDirectory() as folder:
                a = FakeACS(Path(folder))
                a.device = {'_deviceId': {'_SerialNumber': 'fixture'}}

                def api(path):
                    if path == '/presets':
                        return (200, [{'_id': 'apply_test', 'precondition': 'DeviceID.SerialNumber = "fixture"', 'configurations': [{'name': 'force'}]}])
                    return (200, [{'_id': 'force', 'script': 'declare("' + tol.INTERVAL + '", {value: 1}, {value: 30});'}])
                a.api = api
                a.value = lambda p: 30
                s = tol.TOLSuite(a, True)
                self.assertEqual(s.interval_policy(), (30, 'apply_test/force'))
                with self.assertRaises(tol.TestBlocked):
                    s.periodic_600()
                self.assertEqual(s.results[10]['status'], 'BLOCCATO')

        def test_unreachable_ping_is_blocked_and_settings_restored(self):
            with tempfile.TemporaryDirectory() as folder:
                a = FakeACS(Path(folder))
                root = 'Device.IP.Diagnostics.IPPing'
                values = {root + '.DiagnosticsState': 'Completed', root + '.SuccessCount': 0, root + '.FailureCount': 3}
                a.value = lambda p: values.get(p)
                a.node = lambda p: {'_value': values.get(p)}
                a.change = lambda changes: None
                a.refresh = lambda paths: None
                restored = []
                a.restore = lambda paths: restored.extend(paths)
                s = tol.TOLSuite(a, True)
                s.diagnostic_target = lambda: 'http://192.0.2.1/test_50MB.bin'
                s.diagnostic(30, 'IPPing')
                self.assertEqual(s.results[30]['status'], 'BLOCCATO')
                self.assertEqual(s.results[30]['values'][root + '.FailureCount'], '3')
                self.assertIn(root + '.Host', restored)

        def test_64_bit_counters_from_acs_are_exact(self):
            self.assertEqual(tol.counter_value({'_type': 'xsd:unsignedLong', '_value': str(2 ** 64 - 1)}), 2 ** 64 - 1)
            self.assertEqual(tol.counter_value({'_type': 'xsd:unsignedInt', '_value': 0}), 0)
            for value in ['-1', '1.5', str(2 ** 64), True]:
                self.assertIsNone(tol.counter_value({'_type': 'xsd:unsignedLong', '_value': value}))

        def test_restore_failure_stops_later_writes(self):
            a = tol.ACS.__new__(tol.ACS)
            a.protection = {'Device.Test': (1, 'xsd:int')}
            a.restore_failed = False
            a.event = lambda *args, **kwargs: None

            def failed(*args, **kwargs):
                raise tol.CWMPFault('fixture')
            a.change = failed
            with self.assertRaises(tol.CWMPFault):
                a.restore(['Device.Test'])
            self.assertTrue(a.restore_failed)
            self.assertIn('Device.Test', a.protection)

    def packet_fields(interval=30, event='2 PERIODIC', retry=0, responses=True):
        lines=[]
        for i in range(3):
            body=(f'<e:Envelope xmlns:e="http://schemas.xmlsoap.org/soap/envelope/" xmlns:c="urn:dslforum-org:cwmp-1-2"><e:Header><c:ID>unit-{i}</c:ID></e:Header><e:Body><c:Inform><DeviceId><SerialNumber>UNIT-CPE</SerialNumber></DeviceId><Event><EventStruct><EventCode>{event}</EventCode></EventStruct></Event><RetryCount>{retry}</RetryCount></c:Inform></e:Body></e:Envelope>').encode().hex()
            ack=(f'<e:Envelope xmlns:e="http://schemas.xmlsoap.org/soap/envelope/" xmlns:c="urn:dslforum-org:cwmp-1-2"><e:Header><c:ID>unit-{i}</c:ID></e:Header><e:Body><c:InformResponse/></e:Body></e:Envelope>').encode().hex()
            stamp=utcnow().timestamp()-90+i*interval
            for nat in range(2):
                stream=i*2+nat;frame=i*4+nat+1
                lines.append(f'{frame}|{stamp+nat*.00001}|{stream}|POST||{body}')
                if responses:lines.append(f'{frame+2}|{stamp+.1+nat*.00001}|{stream}||200|{ack}')
        return '\n'.join(lines)

    class PipelineTests(unittest.TestCase):
        def rows(self,**kwargs):
            with patch.object(tol,'_command',return_value=SimpleNamespace(stdout=packet_fields(**kwargs))):
                return pcap_periodic_rows('unused.pcap',7557,'UNIT-CPE')
        def test_wire_inform_response_and_nat_deduplication(self):
            rows=self.rows()
            self.assertEqual(len(rows),3)
            self.assertTrue(_pcap_measure(rows,30)['passed'])
            self.assertTrue(all(r['response_frame'] for r in rows))
        def test_bootstrap_is_not_periodic(self):
            self.assertEqual(self.rows(event='0 BOOTSTRAP'),[])
        def test_wrong_cpe_is_excluded(self):
            with patch.object(tol,'_command',return_value=SimpleNamespace(stdout=packet_fields())):
                self.assertEqual(pcap_periodic_rows('unused',7557,'OTHER-CPE'),[])
        def test_retry_does_not_pass(self):
            self.assertFalse(_pcap_measure(self.rows(retry=1),30)['passed'])
        def test_missing_inform_response_does_not_pass(self):
            self.assertFalse(_pcap_measure(self.rows(responses=False),30)['passed'])
        def test_wrong_interval_does_not_pass(self):
            self.assertFalse(_pcap_measure(self.rows(interval=60),30)['passed'])
        def test_invalid_xml_is_ignored(self):
            with patch.object(tol,'_command',return_value=SimpleNamespace(stdout='1|1|0|POST||ff')):
                self.assertEqual(decode_pcap_messages('unused',7557),[])
        def test_old_packets_cannot_be_made_current_by_new_summary(self):
            with tempfile.TemporaryDirectory() as tmp:
                base=Path(tmp);folder=base/'logs'/'unit'/'periodic';folder.mkdir(parents=True)
                pcap=folder/'proof.pcap';pcap.write_bytes(b'unit')
                now=utcnow();device={'_id':'UNIT','_deviceId':{'_SerialNumber':'UNIT-CPE'},'Device':{'ManagementServer':{
                    'PeriodicInformInterval':{'_value':30},'PeriodicInformEnable':{'_value':True},'PeriodicInformTime':{'_value':'2026-01-01T00:00:00Z'}}}}
                proof={'evidence_type':'pcap','device_sha256':hashlib.sha256(b'UNIT').hexdigest(),'ended':now.isoformat(),
                       'expected_interval_s':30,'final_interval_s':30,'periodic_inform_time':'2026-01-01T00:00:00Z',
                       'periodic_inform_enable':True,'pcap':str(pcap.relative_to(base)),'pcap_sha256':hashlib.sha256(b'unit').hexdigest(),'cwmp_port':7557}
                (folder/'summary.json').write_text(json.dumps(proof))
                rows=self.rows()
                for row in rows:row['epoch']-=7200
                with patch.object(tol,'BASE',base),patch.object(tol,'pcap_periodic_rows',return_value=rows):
                    with self.assertRaises(TestBlocked):verified_periodic_evidence(device,30)
                pcap.write_bytes(b'changed')
                with patch.object(tol,'BASE',base):
                    with self.assertRaises(TestBlocked):verified_periodic_evidence(device,30)
        def test_datamodel_walk_keeps_children_of_objects(self):
            model={'Device':{'Object':{'_object':True,'Child':{'_object':False,'_value':7}}}}
            self.assertEqual([p for p,_ in datamodel_parameters(model)],['Device.Object.Child'])
        def test_credential_repeated_in_other_field_is_redacted(self):
            parameters=[('Device.WiFi.KeyPassphrase',{'_value':'unit-password'}),('Device.DeviceInfo.SerialNumber',{'_value':'prefix-unit-password'})]
            secrets=credential_values(parameters)
            self.assertEqual(display_value(*parameters[0]),'[oscurato]')
            self.assertNotIn('unit-password',redact_secrets(parameters[1][1]['_value'],secrets))
        def test_json_masks_credentials_repeated_in_metadata(self):
            with tempfile.TemporaryDirectory() as tmp:
                device={'_id':'prefix-12345678','_deviceId':{'_SerialNumber':'unit-12345678'},'Device':{'WiFi':{'KeyPassphrase':{'_value':'12345678','_object':False}},'DeviceInfo':{'SerialNumber':{'_value':991234567899,'_object':False}}}}
                with patch.object(tol,'BASE',Path(tmp)),contextlib.redirect_stdout(io.StringIO()):export_datamodel(device)
                content=(Path(tmp)/'cpe_datamodel_full.json').read_text()
                self.assertNotIn('12345678',content)
                self.assertEqual(len(datamodel_parameters(json.loads(content))),2)
        def test_expected_300_actual_30_fails(self):
            with tempfile.TemporaryDirectory() as tmp:
                target=Path(tmp)/'audit.xlsx';wb=openpyxl.Workbook();ws=wb.active;ws.title='Sheet2';ws.append(['Parametro','Atteso']);ws.append([INTERVAL,300]);wb.save(target)
                device={'_id':'UNIT','Device':{'ManagementServer':{'PeriodicInformInterval':{'_value':30}}}}
                with patch.object(tol,'EXCEL_FILE',str(target)),contextlib.redirect_stdout(io.StringIO()):run_test_suite(device)
                self.assertEqual(openpyxl.load_workbook(target)['Sheet2']['D2'].value,'FAIL')
        def test_report_only_never_runs_active_suite(self):
            with patch.object(tol,'fetch_device',return_value={'_id':'UNIT'}),patch.object(tol,'export_datamodel'),patch.object(tol,'run_test_suite'),patch.object(tol,'build_presentation'),patch.object(tol,'verified_periodic_evidence',return_value={'passed':True}),patch.object(tol,'run_tol') as live:
                self.assertEqual(main(['--report-only']),0);live.assert_not_called()
        def test_template_report_keeps_editorial_and_has_no_workbooks(self):
            from zipfile import ZipFile
            template=BASE/'report_audit_tr069_guida.pptx'
            if not template.is_file():template=Path(REPORT_TEMPLATE)
            if not template.is_file():self.skipTest('Modello PPTX non disponibile')
            original=Presentation(str(template))
            device={'_id':'offline-unit-device','Device':{'RootDataModelVersion':{'_object':False,'_value':'2.13','_type':'xsd:string'}}}
            with tempfile.TemporaryDirectory() as tmp:
                output=Path(tmp)/'report.pptx'
                with patch.object(tol,'REPORT_TEMPLATE',template),patch.object(tol,'OUTPUT_PPTX',str(output)),patch.object(tol,'PERIODIC_LOG_PATH',str(Path(tmp)/'missing.log')),patch.object(tol,'report_tol_snapshot',return_value=(78,collections.Counter({'PASS':78}),0,'offline')),patch.object(tol,'verified_periodic_evidence',side_effect=TestBlocked('Nessuna cattura nel test isolato')),contextlib.redirect_stdout(io.StringIO()):
                    build_presentation(device)
                result=Presentation(str(output))
                self.assertGreater(len(result.slides),13)
                for i in range(8,13):self.assertEqual(original.slides[i]._element.xml,result.slides[i]._element.xml)
                with ZipFile(output) as package:
                    self.assertFalse(any(n.startswith('ppt/embeddings/') for n in package.namelist()))
                tol_table=template_shape(result.slides[2],7).table
                self.assertEqual(tol_table.cell(1,1).text,'78')
                self.assertIn('WARN',template_shape(result.slides[7],5).text)
                paths=[row.cells[0].text.replace('\n','') for slide in list(result.slides)[13:] for shape in slide.shapes if shape.has_table for row in list(shape.table.rows)[1:]]
                self.assertIn('Device.RootDataModelVersion',paths)

        def test_capture_cleanup_removes_only_its_helper(self):
            with tempfile.TemporaryDirectory() as tmp:
                cap=PeriodicCapture(None,Path(tmp),60,7557);cap.container='owned-unit-helper';commands=[]
                def fake(args,timeout=20):
                    commands.append(args)
                    if args[1]=='inspect':return SimpleNamespace(stdout='[{"State":{"Running":false}}]',stderr='')
                    if args[1]=='cp':raise subprocess.CalledProcessError(1,args)
                    return SimpleNamespace(stdout='',stderr='')
                with patch.object(tol,'_command',side_effect=fake),patch.object(subprocess,'run'):
                    with self.assertRaises(subprocess.CalledProcessError):cap.close()
                self.assertIn(['docker','rm','owned-unit-helper'],commands)
                self.assertIsNone(cap.container)

    suite=unittest.TestSuite([unittest.defaultTestLoader.loadTestsFromTestCase(TOLTests),unittest.defaultTestLoader.loadTestsFromTestCase(PipelineTests)])
    result=unittest.TextTestRunner(verbosity=2).run(suite)
    return 0 if result.wasSuccessful() else 1


if __name__=='__main__':
    raise SystemExit(main())
