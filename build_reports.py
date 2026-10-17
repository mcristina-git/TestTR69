import json
import os
import subprocess
import sys

for pkg in ["openpyxl", "python-pptx"]:
    try:
        __import__(pkg)
    except ImportError:
        subprocess.check_call([sys.executable, "-m", "pip", "install", pkg])

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

with open('/home/cristina/cpe-testbed/all_cpe_params.json') as f:
    all_params = json.load(f)

total_count = len(all_params)
print(f"Elaborazione di {total_count} parametri in corso...")

# ----------------- 1. EXCEL -----------------
wb = Workbook()
ws = wb.active
ws.title = "TR-181 Full DataModel"
ws.views.sheetView[0].showGridLines = True

headers = ["#", "Parametro TR-181 / TR-069", "Valore Attuale", "Ramo"]
ws.append(headers)

hdr_fill = PatternFill(start_color="102542", end_color="102542", fill_type="solid")
hdr_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
thin_border = Border(
    left=Side(style="thin", color="D0D7DE"),
    right=Side(style="thin", color="D0D7DE"),
    top=Side(style="thin", color="D0D7DE"),
    bottom=Side(style="thin", color="D0D7DE")
)

for col in range(1, 5):
    c = ws.cell(row=1, column=col)
    c.fill = hdr_fill
    c.font = hdr_font
    c.alignment = Alignment(horizontal="center", vertical="center")

for idx, (path, val) in enumerate(sorted(all_params.items()), start=1):
    parts = path.split('.')
    branch = ".".join(parts[:2]) if len(parts) > 2 else parts[0]
    ws.append([idx, path, val, branch])
    r_idx = idx + 1
    row_fill = PatternFill(start_color="F6F8FA" if idx % 2 == 0 else "FFFFFF", fill_type="solid")
    for c_idx in range(1, 5):
        c = ws.cell(row=r_idx, column=c_idx)
        c.fill = row_fill
        c.border = thin_border
        c.font = Font(name="Calibri", size=10)
        if c_idx == 1:
            c.alignment = Alignment(horizontal="center")

ws.column_dimensions['A'].width = 8
ws.column_dimensions['B'].width = 65
ws.column_dimensions['C'].width = 40
ws.column_dimensions['D'].width = 25
ws.freeze_panes = "A2"

excel_path = "/home/cristina/cpe-testbed/Report_Parametri_CPE_TCL.xlsx"
wb.save(excel_path)
print(f"[+] File Excel completo generato: {excel_path}")

# ----------------- 2. POWERPOINT -----------------
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

DARK_BLUE = RGBColor(16, 37, 66)
TEXT_MUTED = RGBColor(100, 116, 139)

# Slide 1: Cover
s1 = prs.slides.add_slide(blank_layout)
tb1 = s1.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.3), Inches(3.0))
p = tb1.text_frame.paragraphs[0]
p.text = "CPE 5G TCL: Censimento Completo Data Model TR-181"
p.font.name = "Arial"
p.font.size = Pt(30)
p.font.bold = True
p.font.color.rgb = DARK_BLUE
p_sub = tb1.text_frame.add_paragraph()
p_sub.text = f"Testbed GenieACS CWMP - Totale parametri esportati: {total_count}"
p_sub.font.size = Pt(16)
p_sub.font.color.rgb = TEXT_MUTED

# Slide tabelle a lotti di 16 parametri per pagina
items = sorted(all_params.items())
page_size = 16
pages = (total_count + page_size - 1) // page_size

for p_num in range(pages):
    chunk = items[p_num * page_size : (p_num + 1) * page_size]
    s = prs.slides.add_slide(blank_layout)
    
    # Intestazione slide
    tb = s.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.1))
    p = tb.text_frame.paragraphs[0]
    p.text = f"Data Model TR-181 - Parametri Rilevati ({p_num + 1}/{pages})"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = DARK_BLUE
    p2 = tb.text_frame.add_paragraph()
    p2.text = f"Parametri da {p_num * page_size + 1} a {min((p_num + 1) * page_size, total_count)} di {total_count}"
    p2.font.size = Pt(12)
    p2.font.color.rgb = TEXT_MUTED
    
    # Tabella
    t_shape = s.shapes.add_table(len(chunk) + 1, 2, Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2))
    table = t_shape.table
    table.columns[0].width = Inches(7.5)
    table.columns[1].width = Inches(4.2)
    
    table.cell(0, 0).text = "Parametro TR-181"
    table.cell(0, 1).text = "Valore Registrato"
    for c in range(2):
        cell = table.cell(0, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = DARK_BLUE
        cell.text_frame.paragraphs[0].font.bold = True
        cell.text_frame.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
        cell.text_frame.paragraphs[0].font.size = Pt(11)
        
    for r_idx, (path, val) in enumerate(chunk, start=1):
        table.cell(r_idx, 0).text = path
        table.cell(r_idx, 1).text = str(val)[:50]
        for c in range(2):
            cell_p = table.cell(r_idx, c).text_frame.paragraphs[0]
            cell_p.font.name = "Arial"
            cell_p.font.size = Pt(9.5)

ppt_path = "/home/cristina/cpe-testbed/Report_TR143_TCL_5G.pptx"
prs.save(ppt_path)
print(f"[+] File PowerPoint aggiornato: {ppt_path}")
